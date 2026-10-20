"""Construye el manual ampliado a partir de las 46 capturas aportadas."""
from pathlib import Path
from io import BytesIO
from collections import defaultdict
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import base_visual as B
from contenido_ampliado import DETAILS

HERE=Path(__file__).resolve().parent
PDF=HERE/'Manual METRONET completo - Administrador y Usuario.pdf'
DOCX=HERE/'Manual METRONET completo - Administrador y Usuario.docx'
DETAIL_PDF=HERE/'_detalle.pdf'
by_base=defaultdict(list)
for detail in DETAILS: by_base[detail['base']].append(detail)
assert sorted(by_base)==list(range(1,len(B.PAGES)+1))

DETAIL_SHOTS={
    'Metas y POI: acciones y lugares objetivo':'6.59.24',
    'Red de referencia: edición y prueba privada':'6.59.30',
    'Historial: versiones y limpieza':'6.59.51',
    'Fuentes del mapa y unidades del juego':'7.00.22',
    'Sesión y recuperación de contraseña':'6.57.41',
    'Inicio, menú de cuenta y perfil':'6.58.29',
    'Campaña: elegir nivel y leer la misión':'7.00.44',
    'Tutorial: guía guiada y ayuda contextual':'7.00.57',
    'Mis diseños: recuperar y gestionar':'7.02.08',
    'Reglas: puntos, aprobación y acciones':'6.58.42',
    'UV, UT, ranking y trofeos':'6.59.03',
}

def detail_picture(d):
    key=DETAIL_SHOTS.get(d['title'])
    if not key:return B.FIG/f'main-{d["base"]:02}.png'
    target=B.FIG/f'detail-{d["base"]:02}-{key}.png'
    if not target.exists():
        im=B.crop_img(key,(0,270,1987,1235));im.thumbnail((1400,850));im.save(target,optimize=True)
    return target

def txt(c,s,x,top,w,size=11,bold=False,color=B.INK,leading=None):
    style=ParagraphStyle('p',fontName='Calibri-Bold' if bold else 'Calibri',fontSize=size,leading=leading or size*1.18,textColor=HexColor(color))
    p=Paragraph(s,style);_,h=p.wrap(w,2000);p.drawOn(c,x,top-h);return h

def draw_detail_pdf(c,d,n):
    c.setFillColor(HexColor(B.WHITE));c.rect(0,0,1000,700,fill=1,stroke=0)
    c.setFillColor(HexColor(B.NAVY));c.rect(0,678,1000,22,fill=1,stroke=0)
    part='ADMINISTRADOR' if d['base']<=10 else 'USUARIO'
    txt(c,'METRONET  /  MANUAL COMPLETO',42,661,600,11,True,B.BLUE)
    txt(c,'PARTE 1 · ADMINISTRADOR' if part=='ADMINISTRADOR' else 'PARTE 2 · USUARIO',740,661,220,10,True,B.MUTED)
    txt(c,d['title'],42,629,905,27,True,B.NAVY)
    c.setStrokeColor(HexColor(B.SKY));c.setLineWidth(2);c.line(42,579,958,579)
    txt(c,'PARA QUÉ SIRVE',42,562,520,11,True,B.BLUE)
    txt(c,d['purpose'],42,540,520,12,False,B.INK,14.4)
    txt(c,'DÓNDE ENTRAR',42,459,520,11,True,B.BLUE)
    txt(c,d['route'],42,438,520,11.5,False,B.INK,13.4)
    img=Image.open(detail_picture(d)).convert('RGB')
    iw,ih=img.size; sc=min(365/iw,190/ih);w=iw*sc;h=ih*sc
    c.drawImage(ImageReader(img),595,567-h,width=w,height=h)
    c.setStrokeColor(HexColor(B.LINE));c.rect(595,567-h,w,h,fill=0,stroke=1)
    txt(c,'CONTROLES Y ELEMENTOS',42,382,916,12,True,B.BLUE)
    mid=(len(d['controls'])+1)//2
    for col,items in enumerate((d['controls'][:mid],d['controls'][mid:])):
        y=357;x=42+col*462
        for s in items:
            c.setFillColor(HexColor(B.SKY));c.circle(x+4,y-6,2.5,fill=1,stroke=0)
            h=txt(c,s,x+13,y,435,10.4,False,B.INK,12.2)
            y-=max(h+11,32)
        if y<185: print('AVISO control largo',d['title'],col,y)
    c.setStrokeColor(HexColor(B.LINE));c.line(42,174,958,174)
    txt(c,'PASOS',42,162,916,11,True,B.BLUE)
    y=143
    for i,s in enumerate(d['steps'],1):
        h=txt(c,f'<b>{i}.</b> {s}',44,y,900,10.2,False,B.INK,11.6)
        y-=h+5
    if y<47: print('AVISO pasos largos',d['title'],y)
    c.setFillColor(HexColor(B.PALE));c.roundRect(42,29,916,43,4,fill=1,stroke=0)
    txt(c,f'<b>Resultado:</b> {d["result"]}',52,66,894,9.6,False,B.DEEP,11)
    if d['caution']:
        # Nota en el archivo editable y matriz; en PDF figura debajo de la ruta.
        txt(c,f'<b>Atención:</b> {d["caution"]}',42,412,520,9.5,False,B.MUTED,10.8)
    c.setFont('Calibri',9);c.setFillColor(HexColor(B.MUTED));c.drawRightString(958,12,f'GUÍA DETALLADA · {n}')
    c.showPage()

def make_detail_pdf():
    c=canvas.Canvas(str(DETAIL_PDF),pagesize=(1000,700),pageCompression=1)
    n=0
    for base in range(1,len(B.PAGES)+1):
        for d in by_base[base]:
            n+=1;draw_detail_pdf(c,d,n)
    c.save()

def merge_pdf():
    base=PdfReader(str(B.PDF)); detail=PdfReader(str(DETAIL_PDF)); writer=PdfWriter()
    writer.add_page(base.pages[0]); j=0
    admin=None;usuario=None
    for i in range(1,len(base.pages)):
        page_index=len(writer.pages)
        writer.add_page(base.pages[i])
        if i<=10:
            if admin is None:admin=writer.add_outline_item('Parte 1 · Administrador',page_index)
            parent=admin
        else:
            if usuario is None:usuario=writer.add_outline_item('Parte 2 · Usuario',page_index)
            parent=usuario
        writer.add_outline_item(B.PAGES[i-1]['title'],page_index,parent=parent)
        for _ in by_base[i]:
            writer.add_page(detail.pages[j]);j+=1
    assert j==len(detail.pages)
    writer.add_metadata({'/Title':'Manual METRONET completo — Administrador y Usuario','/Subject':'Guía visual y operativa de METRONET'})
    with PDF.open('wb') as f:writer.write(f)

def rgb(h):
    h=h.lstrip('#');return RGBColor(*(int(h[i:i+2],16) for i in (0,2,4)))

def add_p(d,text='',style=None,size=None,bold=False,color=None,after=3):
    p=d.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(after)
    if text:
        r=p.add_run(text);r.font.name='Calibri'
        if size:r.font.size=Pt(size)
        r.bold=bold
        if color:r.font.color.rgb=rgb(color)
    return p

def picture(paragraph,path,max_w,max_h):
    im=Image.open(path); w,h=im.size;scale=min(max_w/w,max_h/h)
    paragraph.add_run().add_picture(str(path),width=Inches(w*scale),height=Inches(h*scale))

def doc_page_header(doc,part,title,detail=False):
    p=add_p(doc,'PARTE 1 · ADMINISTRADOR' if part=='ADMINISTRADOR' else 'PARTE 2 · USUARIO',size=9,bold=True,color=B.BLUE,after=0)
    p.paragraph_format.page_break_before=True
    add_p(doc,title,style='Heading 1',after=4)
    if detail:add_p(doc,'GUÍA DETALLADA',size=9,bold=True,color=B.BLUE,after=3)

def overview_doc(doc,idx,item):
    doc_page_header(doc,item['part'],item['title'])
    add_p(doc,item['lead'],size=10.5,color=B.MUTED,after=5)
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    table.columns[0].width=Inches(6.5);table.columns[1].width=Inches(4.25)
    a,b=table.rows[0].cells;a.width=Inches(6.5);b.width=Inches(4.25)
    picture(a.paragraphs[0],B.FIG/f'main-{idx:02}.png',6.45,3.6)
    b.text=''
    for n,(label,description) in enumerate(item['items'],1):
        p=b.add_paragraph();p.paragraph_format.space_after=Pt(7)
        r=p.add_run(f'{n}  {label}\n');r.bold=True;r.font.name='Calibri';r.font.size=Pt(11);r.font.color.rgb=rgb(B.BLUE)
        r=p.add_run(description);r.font.name='Calibri';r.font.size=Pt(9.5)
    add_p(doc,'PANTALLAS RELACIONADAS',size=9.5,bold=True,color=B.BLUE,after=2)
    n=len(item['extra']);t=doc.add_table(rows=1,cols=n);t.autofit=False
    for j,label in enumerate(item['extra_names']):
        cell=t.cell(0,j);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        picture(cell.paragraphs[0],B.FIG/f'extra-{idx:02}-{j}.png',min(3.4,10.6/n),1.35)
        p=cell.add_paragraph();p.paragraph_format.space_after=Pt(0)
        r=p.add_run(label);r.font.name='Calibri';r.font.size=Pt(8.5);r.bold=True;r.font.color.rgb=rgb(B.DEEP)
    if item.get('note'):add_p(doc,item['note'],size=9,color=B.MUTED,after=0)

def detail_doc(doc,d):
    part='ADMINISTRADOR' if d['base']<=10 else 'USUARIO'
    doc_page_header(doc,part,d['title'],True)
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    table.columns[0].width=Inches(6.8);table.columns[1].width=Inches(3.9)
    left,right=table.rows[0].cells;left.width=Inches(6.8);right.width=Inches(3.9)
    left.text='';right.text=''
    p=left.paragraphs[0];p.paragraph_format.space_after=Pt(3)
    r=p.add_run('PARA QUÉ SIRVE\n');r.bold=True;r.font.name='Calibri';r.font.size=Pt(10);r.font.color.rgb=rgb(B.BLUE)
    r=p.add_run(d['purpose']);r.font.name='Calibri';r.font.size=Pt(10)
    p=left.add_paragraph();p.paragraph_format.space_after=Pt(0)
    r=p.add_run('DÓNDE ENTRAR\n');r.bold=True;r.font.name='Calibri';r.font.size=Pt(10);r.font.color.rgb=rgb(B.BLUE)
    r=p.add_run(d['route']);r.font.name='Calibri';r.font.size=Pt(10)
    picture(right.paragraphs[0],detail_picture(d),3.75,1.65)
    add_p(doc,'CONTROLES Y ELEMENTOS',size=10,bold=True,color=B.BLUE,after=2)
    mid=(len(d['controls'])+1)//2
    ct=doc.add_table(rows=1,cols=2);ct.autofit=False
    for col,items in enumerate((d['controls'][:mid],d['controls'][mid:])):
        cell=ct.cell(0,col);cell.text=''
        for s in items:
            p=cell.add_paragraph(style='Normal');p.paragraph_format.space_after=Pt(4)
            r=p.add_run('• '+s);r.font.name='Calibri';r.font.size=Pt(9.4)
    add_p(doc,'PASOS',size=10,bold=True,color=B.BLUE,after=2)
    for i,s in enumerate(d['steps'],1):add_p(doc,f'{i}. {s}',size=9.5,after=2)
    add_p(doc,'RESULTADO  '+d['result'],size=9.5,bold=False,color=B.DEEP,after=2)
    if d['caution']:add_p(doc,'ATENCIÓN  '+d['caution'],size=9,color=B.MUTED,after=0)

def make_docx():
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(12);sec.page_height=Inches(8.4);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.top_margin=Inches(.37);sec.bottom_margin=Inches(.3);sec.left_margin=Inches(.55);sec.right_margin=Inches(.55)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(9.5);normal.font.color.rgb=rgb(B.INK);normal.paragraph_format.space_after=Pt(2)
    h=doc.styles['Heading 1'];h.font.name='Calibri';h.font.bold=True;h.font.size=Pt(22);h.font.color.rgb=rgb(B.NAVY)
    for style in ('Normal','Heading 1'):doc.styles[style]._element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    p=add_p(doc);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;picture(p,B.LOGO,4.1,3.2)
    p=add_p(doc,'MANUAL METRONET',size=30,bold=True,color=B.NAVY,after=10);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=add_p(doc,'Parte 1 · Administrador     /     Parte 2 · Usuario',size=18,color=B.BLUE,after=8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=add_p(doc,'Guía visual y operativa · 46 capturas del juego',size=13,color=B.MUTED);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for idx,item in enumerate(B.PAGES,1):
        overview_doc(doc,idx,item)
        for d in by_base[idx]:detail_doc(doc,d)
    doc.core_properties.title='Manual METRONET completo — Administrador y Usuario'
    doc.save(DOCX)

def main():
    B.main()  # genera las figuras recortadas, con referencias, y el PDF base
    make_detail_pdf();merge_pdf();make_docx()
    print(PDF);print(DOCX);print('paginas_pdf',len(PdfReader(str(PDF)).pages),'capturas',len(set(B.used)))

if __name__=='__main__':main()

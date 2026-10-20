from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
OUT = BASE / 'Manual METRONET Usuario y Administrador 2026-10-18 revision 3.docx'
M = BASE / 'materiales'

d = Document()
s = d.sections[0]
s.top_margin = Inches(.74); s.bottom_margin = Inches(.7)
s.left_margin = Inches(.83); s.right_margin = Inches(.83)
s.header_distance = Inches(.3); s.footer_distance = Inches(.35)

styles = d.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor(28, 34, 43)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.15
for name, size, before, after in [('Title', 24, 0, 12), ('Heading 1', 15, 16, 8), ('Heading 2', 11.5, 11, 5), ('Heading 3', 10.5, 8, 4)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name in ('Title','Heading 1') else 'Aptos'
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
styles['Caption'].font.name = 'Aptos'; styles['Caption'].font.size = Pt(8.5)
styles['Caption'].font.color.rgb = RGBColor(70, 78, 89)
styles['Caption'].paragraph_format.space_after = Pt(9)
title_style_ppr = styles['Title']._element.get_or_add_pPr()
border = title_style_ppr.find(qn('w:pBdr'))
if border is not None: title_style_ppr.remove(border)

footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('METRONET  ·  Manual de usuario  ·  18.10.2026  ·  ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def title(t): d.add_paragraph(t, 'Title')
def h1(t): d.add_heading(t, 1)
def h2(t): d.add_heading(t, 2)
def p(t): d.add_paragraph(t)
def steps(items):
    for i, item in enumerate(items,1):
        q=d.add_paragraph(style='Normal'); q.paragraph_format.left_indent=Inches(.22); q.paragraph_format.first_line_indent=Inches(-.22)
        r=q.add_run(f'{i}. '); r.bold=True; q.add_run(item)
def bullet(items):
    for item in items:
        q=d.add_paragraph(item, style='List Bullet'); q.paragraph_format.space_after=Pt(3)
def outcome(t):
    q=d.add_paragraph(); q.add_run('Resultado esperado. ').bold=True; q.add_run(t)
capture_count=0
def capture(n, screen):
    global capture_count
    capture_count += 1
    q=d.add_paragraph(style='Caption')
    q.paragraph_format.keep_with_next=True
    r=q.add_run(f'Captura pendiente {capture_count:02d}. ')
    r.bold=True
    q.add_run(screen)
def fig(filename, caption, width=4.65):
    img=M/filename
    if img.exists():
        q=d.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.keep_with_next=True
        shape=q.add_run().add_picture(str(img), width=Inches(width))
        shape._inline.docPr.set('descr', caption)
        c=d.add_paragraph(caption, 'Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style='Table Grid'
    for i, x in enumerate(headers): t.rows[0].cells[i].text=x
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row): c[i].text=str(x)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    trPr=t.rows[0]._tr.get_or_add_trPr(); header=OxmlElement('w:tblHeader'); header.set(qn('w:val'),'true'); trPr.append(header)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.font.size=Pt(8.5)
            tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'), '1E3656' if ri==0 else ('F1F5F8' if ri%2==0 else 'FFFFFF')); tcPr.append(sh)
            if ri==0:
                for para in cell.paragraphs:
                    for run in para.runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(1)

title('Manual de usuario METRONET')
p('Guía práctica para estudiantes de logística y personas administradoras. La Parte A acompaña el recorrido desde el acceso hasta el resultado de una red. La Parte B explica los siete módulos exclusivos de Administración, sus controles y sus efectos.')
fig('logo.png','Figura 1. Logotipo original de METRONET incluido en la aplicación.',3.1)
p('Versión local comprobada: b5168659bebe276fee89f37963b97daf22992503. Revisión del manual: 3 · 18 de octubre de 2026.')
p('Alcance de la evidencia. IMPLEMENTADO identifica funciones presentes en la versión local. Acceso, navegación, edición y administración fueron observados en la aplicación; la red del intento de prueba estaba vacía y no permitió abrir Simulación. Sus controles se contrastaron con la implementación, pero aún falta observarlos durante una ejecución. Los lugares de las capturas actuales están señalados en el texto para incorporarlas en una revisión posterior.')

d.add_page_break()
h1('Índice')
for line in ['PARTE A  USUARIO','1 Registro, acceso, recuperación y perfil','2 Campaña y niveles','3 Tutoriales y Extras','4 Edición de la red','5 Mapa, capas, objetivos y POI','6 Guardado, simulación y UV y UT','7 Resultados, diseños, aprendizaje y reglas','8 Ranking, trofeos y Modo Libre','PARTE B  ADMINISTRADOR','9 Ingreso y navegación de Administración','10 Usuarios','11 Diseños y actividad','12 Experiencia de juego','13 Fuentes, unidades y configuración','14 Problemas habituales','15 Glosario']:
    p(line)

h1('PARTE A  USUARIO')
p('Los capítulos 1 a 8 describen la experiencia de una cuenta jugadora. El administrador también puede consultar varias de estas pantallas, pero su cuenta tiene accesos adicionales y no participa de la clasificación como jugador.')
h1('1 Registro, acceso, recuperación y perfil')
h2('Crear una cuenta de estudiante')
capture(1,'Registro: formulario completo y casilla de privacidad, sin datos personales visibles.')
steps(['Desde Iniciar sesión elegí Crear cuenta.','Completá Nombre, Apellido, Correo electrónico y Contraseña. La clave necesita al menos seis caracteres, una mayúscula y un carácter especial; Mostrar permite verla mientras escribís.','Leé Leer el uso de mis datos y aceptá la casilla de privacidad si estás de acuerdo.','Pulsá Crear cuenta. Después ingresá por el acceso de estudiante.'])
outcome('Se crea una cuenta de jugador cuando los datos son válidos y el correo está disponible. Crear cuenta no habilita funciones de administrador.')
h2('Entrar como estudiante')
capture(2,'Ingreso de estudiante: correo, contraseña y accesos a registro y recuperación.')
steps(['Abrí la dirección local que usa el proyecto, normalmente http://127.0.0.1:5173/login.html, con frontend y backend ya iniciados.','Escribí tu correo electrónico y contraseña. El botón Mostrar permite comprobar lo escrito.','Elegí Iniciar sesión. La aplicación conduce a Inicio o al destino protegido desde el que llegaste.'])
outcome('Se ve el inicio con acceso al nivel actual y al Modo Libre si está habilitado.')
h2('Recuperar el acceso')
capture(3,'Recuperar acceso: correo y envío de código; complementar con Verificar código y Nueva contraseña cuando el usuario aporte imágenes.')
capture(0,'Verificar código: seis dígitos y opción de reenvío habilitada luego de la espera visible.')
capture(0,'Nueva contraseña: campos de clave y confirmación, sin mostrar ninguna contraseña.')
steps(['En Iniciar sesión elegí ¿Olvidaste tu contraseña?','En Recuperar acceso escribí el correo de tu cuenta y pulsá Enviar código. La pantalla evita revelar si ese correo existe.','Si recibís el mensaje, abrí Verificar código, ingresá sus seis dígitos dentro del plazo indicado y pulsá Verificar código. Reenviar código se habilita tras la espera visible.','En Nueva contraseña escribí y repetí una clave válida; elegí Cambiar contraseña. Volvé al ingreso.'])
outcome('Tras una verificación válida, la contraseña nueva reemplaza a la anterior. Las dos últimas pantallas sólo aparecen si se inició y verificó una solicitud real; este manual no envió correos ni cambió credenciales.')
h2('Orientarse en la aplicación')
capture(0,'Inicio: tarjetas del nivel actual y de Modo Libre, menú global y menú de usuario.')
p('Para el estudiante, la navegación ofrece Inicio, Niveles, Aprendizaje, Reglas, Ranking y Mis diseños. El administrador ve además Administración. El menú del nombre abre Mi perfil y Cerrar sesión; el control Música cambia el sonido. Durante un nivel se indica la secuencia Nivel → Edición → Simulación → Resultados. Volvé a Inicio para retomar el recorrido que corresponda a tu cuenta.')
p('En Inicio, la tarjeta principal identifica el nivel para continuar y muestra el estado de la campaña. La segunda tarjeta permite entrar a Modo Libre cuando se desbloquea. Si necesitás consultar o cambiar tus datos, abrí Mi perfil desde el menú de usuario; allí se separan datos personales, correo y contraseña.')
h2('Mi perfil')
capture(4,'Mi perfil: secciones de datos personales, correo y contraseña, con información privada ocultada.')
p('Mi perfil muestra rol y fecha de ingreso. Datos personales permite actualizar nombre y apellido; Correo electrónico muestra el correo actual y recibe uno nuevo; Contraseña solicita clave actual, nueva y confirmación, con botones Mostrar y los requisitos visibles. Cada bloque tiene su propio Guardar. Volver a Inicio sale del perfil. Cambiar estos datos modifica la cuenta, por eso el recorrido de verificación sólo observó los formularios.')

h1('2 Campaña y niveles')
h2('Elegir y continuar un nivel')
capture(5,'Niveles: selector numerado, estado y detalle de un nivel disponible.')
capture(0,'Preparación de nivel: relato, metas y botón Iniciar nivel o Continuar.')
steps(['Entrá en Niveles desde el menú.','Seleccioná un número del 01 al 10. La tarjeta muestra nombre, estado, relato y desafío.','Leé el objetivo antes de iniciar o retomar. Elegí el botón de la tarjeta cuando el nivel esté desbloqueado.','La preparación muestra el relato y una animación de viaje. Podés esperar o pulsar Iniciar nivel. Cuando aparece una tarjeta educativa, leé la idea clave y pulsá Continuar para abrir Edición.'])
outcome('Se abre una red del intento actual y aparecen las metas que debés cumplir. Un nivel bloqueado no permite iniciarlo.')
p('La pantalla Niveles tiene un selector numerado y una tarjeta de detalle. En el selector leés estados como disponible, en curso, completado o bloqueado. En el detalle aparece el botón Iniciar, Continuar, Repetir o Bloqueado según el estado. La preparación muestra el relato y las metas antes de entrar al plano.')
p('La campaña progresa de forma secuencial. Los diez niveles enseñan de un primer enlace entre estaciones a redes de dos líneas, cobertura de lugares, transbordos y comparación de simulaciones. Completar una simulación no aprueba por sí solo el nivel: hay que volver a Edición y usar Finalizar red.')
p('En la cuenta administradora observada, el catálogo ofrece acceso a los diez niveles y Modo Libre. Eso es una vista del rol de administración: para un jugador la disponibilidad depende de su propio avance y no debe interpretarse como un desbloqueo general.')
table(['Niveles','Qué practicás'],[
('1 a 3','Primer enlace, extensión, velocidad y duración simulada.'),
('4 a 6','Varias unidades, cobertura de POI y continuidad de la línea.'),
('7 a 8','Dos líneas, transbordo y objetivo territorial.'),
('9 a 10','Comparaciones de ejecuciones y desafío integrador.')],[1.2,5.35])
p('Reiniciar recorrido aparece sólo si hay avance en la campaña. Abre una confirmación; usalo deliberadamente porque cambia el recorrido actual, aunque el sistema conserva el histórico según sus verificaciones. Este manual no ejecutó esa acción.')

h1('3 Tutoriales y Extras')
h2('Seguir o volver a abrir una ayuda')
capture(6,'Edición: diálogo de bienvenida y un paso del tutorial sobre la herramienta que enseña.')
steps(['Al abrir Edición por primera vez, elegí Iniciar tutorial si querés la guía paso a paso, o Continuar para construir directamente.','Avanzá leyendo la instrucción y realizá la acción indicada en la interfaz. Los pasos prácticos comprueban acciones reales, no sólo pulsar Siguiente.','Para volver a verla, abrí el botón Tutorial de la esquina superior derecha y elegí Iniciar tutorial. La Guía rápida resume Estación, Conexión, Metro, Guardar, Simular y Finalizar.'])
outcome('El panel de ayuda cambia de paso cuando se cumple la acción esperada. Las ayudas extensas se presentan al introducir herramientas nuevas y quedan disponibles manualmente.')
p('La campaña presenta la construcción en el nivel 1, POI en el 5 y transbordo en el 7. La simulación introduce controles en los niveles correspondientes. Las guías pueden variar según la versión publicada del intento.')
h2('Extras de Modo Libre')
capture(7,'Extras: tarjeta con título, explicación, práctica y controles de avance.')
p('Al entrar por primera vez a Modo Libre, la aplicación puede ofrecer siete tarjetas Extras. Explican calles, población, tráfico, desplazamientos, horarios, cobertura y relación entre diseño y simulación. Podés volver a consultarlas en Aprendizaje → Modo Libre. Su lectura es material didáctico; no modifica la velocidad, la duración ni la red.')
p('Cada Extra muestra un título, qué ofrece la herramienta, cómo ayuda a la red, una práctica y una advertencia de interpretación. Podés avanzar por las siete tarjetas, salir y volver desde la biblioteca.')

h1('4 Edición de la red')
h2('Crear un recorrido básico')
capture(8,'Edición: mapa, metas, barra de herramientas y acciones con una red propia de demostración.')
steps(['Desde un nivel desbloqueado, entrá en Edición y leé Metas del nivel.','Elegí Estación en Construcción y hacé clic en un lugar válido del mapa. Repetí para ubicar otra estación.','Elegí Línea y creá la línea que vas a usar.','Elegí Conexión, seleccioná la línea activa y uní dos estaciones. Para extenderla, conectá estaciones adicionales de la misma línea.','Elegí Metros y hacé clic sobre un tramo inequívoco para asignar una unidad. Si hay líneas superpuestas, elegí la que corresponde.','Pulsá Guardar y comprobá el aviso Diseño guardado.'])
outcome('El panel presenta estaciones, líneas, tramos y metros del diseño; las metas se actualizan de acuerdo con la red confirmada.')
p('La herramienta Seleccionar permite inspeccionar una estación, línea, conexión o metro. Desde el elemento seleccionado podés editar sus datos; Mover y Eliminar se habilitan cuando corresponde. Cancelar operación devuelve al modo normal. Una eliminación abre un diálogo porque puede quitar conexiones u otras partes de la red.')
p('En Edición distinguí tres zonas: el mapa a la izquierda, Metas del nivel a la derecha (y Lugares a conectar cuando la consigna usa POI), y Herramientas debajo de las metas. Construcción reúne Estación, Línea, Conexión y Metros; Edición reúne Seleccionar, Mover, Eliminar y Cancelar. Acciones reúne Guardar, Simular diseño y, para niveles, Finalizar red. El aviso de estado comunica confirmaciones y errores.')
p('En la visita al nivel 1, las cinco metas visibles fueron: ubicar dos estaciones, crear una línea, crear una conexión, asignar una unidad de metro y simular la red con sus UV/UT actuales. Sin ningún elemento creado, Mover, Eliminar y Cancelar estaban deshabilitados; Seleccionar era la herramienta activa. El mapa ofrecía zoom, Ajustar red y tres controles de orientación. El botón Aprender y el botón Tutorial estaban en la cabecera.')
p('Un transbordo requiere que dos líneas compartan una estación real. Dibujar dos vías que se cruzan sin esa estación no cuenta. En los niveles avanzados se exige al menos un metro por línea.')

h1('5 Mapa, capas, objetivos y POI')
h2('Leer el mapa')
capture(9,'Mapa: zoom, orientación y referencias visibles junto a la red.')
p('El mapa de Montevideo muestra barrios, estaciones, líneas y referencias. Los controles + y − cambian el zoom; restaurar vista vuelve al encuadre de referencia. También hay mandos de orientación. La barra geográfica abre referencias y opciones del mapa. Al seleccionar una zona o un elemento, la interfaz muestra datos relacionados.')
h2('Usar POI y comprobar una meta')
capture(10,'Puntos de interés: lista o búsqueda, marcador y estado de la meta relacionada.')
steps(['Abrí Puntos de interés desde el control de referencias del mapa.','Usá Buscar punto de interés para localizar el lugar solicitado. Podés activar categorías de marcadores.','Seleccioná el POI para centrar la vista y observar su referencia.','Ubicá una estación dentro del radio que muestra la meta; conectala por la línea requerida.','Revisá Metas del nivel y Lugares a conectar. Confirmá que el estado cambie de Pendiente a Cubierto o Completado.'])
outcome('El objetivo se marca como cubierto cuando la geometría y las conexiones cumplen la consigna; estar visualmente cerca de la estrella no basta.')
p('En los niveles del recorrido, Palacio Legislativo y Torre de las Comunicaciones aparecen como referencias de cobertura. Sus radios de evaluación son reglas de juego y no equivalen directamente a kilómetros. Leé siempre la consigna publicada de tu intento.')
h2('Capas territoriales')
capture(0,'Modo Libre: panel Análisis del territorio con capas y selector horario.')
p('Modo Libre habilita capas de calles, tráfico histórico, población y desplazamientos, además de cobertura y controles horarios. Activá una capa desde Análisis del territorio, ajustá opacidad cuando exista ese control y compará el mismo trazado con distintas vistas. En Simulación libre hay un panel Capas del mapa. La cobertura dibuja 500 m en línea recta alrededor de cada estación: no calcula accesibilidad peatonal ni demanda real. Tráfico y desplazamientos son referencias históricas; no representan datos en vivo.')
p('El selector horario cambia la hora de referencia de tráfico y desplazamientos. La población permanece fija. Si una capa falla, el panel muestra un aviso y permite reintentar; podés seguir trabajando con la red.')

h1('6 Guardado, simulación y UV y UT')
capture(11,'Edición: botón Guardar y aviso de estado; una red preparada para Simular diseño.')
h2('De Edición a Simulación')
steps(['En Edición, pulsá Guardar para conservar la red y revisar las metas.','Cuando la red sea consistente y tenga el metro requerido, elegí Simular diseño. La aplicación comprueba y guarda la preparación antes de abrir Simulación.','En Simulación, observá el mapa, la ficha de unidad y los controles. Elegí Iniciar.','Usá Pausar, Reanudar, Detener o Reiniciar si necesitás revisar la ejecución. Al finalizar, consultá el resultado y regresá a Edición para Finalizar red.'])
outcome('Una ejecución terminada queda registrada. Consultar sus resultados no aprueba la campaña; Finalizar red realiza esa comprobación.')
p('Comprobación en vivo: al pulsar Simular diseño con la red vacía del nuevo intento, apareció «Prepará la red para simular» con la indicación de completar objetivos y asignar un metro. No se abrió Simulación ni se registró una ejecución. La descripción de la pantalla y controles que sigue corresponde a la implementación inspeccionada en el código, no a una visita visual de esa pantalla.')
p('La pantalla Simulación tiene un mapa y un panel de instrumentos. En Acciones están Iniciar, Pausar, Reanudar, Detener y Reiniciar. Debajo se ven Estado, Tiempo simulado y el porcentaje de Ejecución. El panel de parámetros contiene Duración · UT y, cuando corresponde, la velocidad UV, el selector de línea o unidad y el ritmo visual. Ampliar mapa oculta temporalmente parte de los instrumentos para ver mejor el recorrido.')
h2('Ajustar una línea')
capture(12,'Simulación: controles, estado, tiempo, UV y UT de una línea durante una prueba real.')
p('UV expresa velocidad: 1 UV = 1 km por UT. UT expresa tiempo simulado: 1 UT = 1 hora del escenario. En el recorrido vigente, cada línea comparte UV y UT entre sus metros. Al asignar el primer metro, comienza con 1 UV. Si elegís UV, la duración de la línea se calcula como distancia ÷ UV; si elegís UT, la velocidad se calcula como distancia ÷ UT. Otra línea conserva sus propios valores.')
steps(['Seleccioná una línea o unidad según el control disponible en el panel. Comprobá qué línea indica la ficha antes de editar.','Cambiá Velocidad · UV o Duración · UT a un valor mayor que cero. Usá los botones + y − o escribí el número.','Elegí el botón de guardar velocidad o Aplicar UT. Observá el otro valor recalculado para la distancia de esa línea.','Volvé a ejecutar para comparar el resultado. Si el nivel pide una comparación, conservá las demás variables que indique su consigna.'])
outcome('El panel muestra UV y UT de la línea seleccionada y la ejecución nueva refleja el ajuste guardado.')
p('Ritmo visual × controla la reproducción en pantalla; no equivale a UV. Una línea más corta puede terminar antes que otra y quedar en su estación terminal. Los valores mostrados pueden redondearse a dos decimales; la aplicación guarda más precisión internamente.')
table(['Acción','Efecto'],[('Guardar','Conserva la red y actualiza metas.'),('Simular','Comprueba, ejecuta y registra una prueba de la red.'),('Finalizar red','Evalúa la consigna y, si corresponde, aprueba el nivel.')],[1.45,5.1])

h1('7 Resultados, diseños, aprendizaje y reglas')
h2('Interpretar el resultado')
capture(13,'Resultados: ejecución terminada, condiciones y desglose de puntuación si el nivel se aprobó.')
p('Simulación muestra estado, tiempo simulado, porcentaje de ejecución, unidad seguida y mensajes de la red. Tras una ejecución se puede consultar el resultado guardado. Al finalizar un nivel aprobado aparece una vista con puntos obtenidos, posible récord personal y desglose del cálculo. Abrí cada ejecución del desglose para conocer descuentos o condiciones pendientes; elegí Continuar para seguir.')
h2('Volver a un diseño')
capture(14,'Mis diseños: pestañas Nivel y Modo Libre, tarjeta y acciones disponibles.')
steps(['Entrá en Mis diseños.','Elegí Diseño de los niveles o Modo Libre. Las pestañas separan ambas clases de red.','En una tarjeta propia, elegí Abrir diseño para editar o Simular diseño si está preparado. Usá Anterior y Siguiente si la lista tiene varias páginas.'])
outcome('Se abre el diseño elegido con su contexto. La opción Simular puede estar deshabilitada si la red aún no está preparada.')
p('Eliminar diseño pide confirmación y no se puede deshacer. En un diseño de nivel también puede quitar partida, puntaje y resultados, con efecto sobre progreso y ranking. No lo uses como forma de volver a editar.')
h2('Consultar Aprendizaje y Reglas')
capture(15,'Aprendizaje: selector de nivel y una tarjeta publicada con fuente.')
capture(16,'Reglas: pestañas de Puntuación, Cómo aprobar, Acciones, Velocidad y tiempo, y Ranking.')
p('Aprendizaje permite elegir un nivel y abrir sus siete tarjetas publicadas cuando está desbloqueado. Cada tarjeta incluye imagen, explicación, idea clave y fuente. La pestaña Modo Libre reúne las siete tarjetas Extras luego de ingresar a ese modo.')
p('Reglas del juego se organiza en Puntuación, Cómo aprobar, Acciones, Velocidad y tiempo, y Ranking. Consultá Cómo aprobar antes de Finalizar red; en Velocidad y tiempo se explican UV y UT. La pestaña Puntuación detalla descuentos y cálculo. Si una regla de una publicación histórica difiere, la vista de resultado puede mostrar la regla original de ese intento.')
p('En Mis diseños, cada tarjeta incluye nombre, estado y acciones para abrir o simular cuando estén disponibles. Los contadores de las pestañas indican cuántos diseños hay en cada grupo. En Aprendizaje, el selector 01 a 10 y Modo Libre cambia las tarjetas visibles. Las pantallas pueden indicar bloqueado o error de carga; Reintentar aparece cuando la consulta falla.')

h1('8 Ranking, trofeos y Modo Libre')
h2('Ranking y reconocimientos')
capture(17,'Ranking: Clasificación o Puntos por nivel, Trofeos y Tu desempeño de una cuenta jugadora.')
p('En Ranking podés alternar Clasificación y Puntos por nivel. Hay navegación entre páginas y secciones Trofeos y Tu desempeño. La mejor marca de un nivel cuenta para la clasificación según la regla visible en Reglas. Los trofeos aparecen cuando se cumplen las condiciones correspondientes; el panel muestra el estado alcanzado para tu cuenta.')
p('Clasificación ordena jugadores y sus puntos; Puntos por nivel desglosa los diez niveles en páginas. Trofeos permite seleccionar un reconocimiento para leer el requisito, incluso cuando todavía está oculto. Tu desempeño resume tus marcas. Si no hay puntuaciones, la pantalla muestra un estado vacío en lugar de una clasificación inventada.')
p('Los reconocimientos visibles son Primera estación, Camino a la cima, Copa de la red, Biblioteca completa, Corona de platino y Premio oculto. Seleccionar una ficha muestra su requisito. En la cuenta administradora observada figuran «En exhibición» y el aviso de que esa vista no atribuye premios; los logros reales corresponden a cuentas jugadoras.')
h2('Diseñar sin una consigna de campaña')
capture(18,'Modo Libre: acceso desbloqueado y editor de una red propia.')
steps(['Completá los diez niveles del recorrido. El acceso a Modo Libre aparecerá disponible en Inicio, Niveles y Mis diseños.','Elegí Modo Libre y creá una red nueva. Si se ofrece la bienvenida de Extras, leé o recorré sus tarjetas.','Ubicá estaciones, construí líneas y conexiones, asigná metros y guardá.','Explorá las capas territoriales y la cobertura. Probá UV o UT por línea, simulá y compará tu decisión con el resultado.'])
outcome('La red queda en la pestaña Modo Libre de Mis diseños. Podés editarla y volver a simularla sin una meta de aprobación de nivel.')
p('Modo Libre sirve para ensayar decisiones de planificación. Las capas no calculan por sí solas demanda, capacidad ni viabilidad real.')

heading9=d.add_heading('PARTE B  ADMINISTRADOR',1)
heading9.paragraph_format.page_break_before=True
p('Esta parte corresponde exclusivamente al rol administrador. Una cuenta jugadora no ve el enlace Administración ni sus siete módulos. Ambos roles pueden consultar pantallas generales, pero un administrador no obtiene marcas ni trofeos de jugador: en Ranking los reconocimientos figuran «En exhibición».')
h1('9 Ingreso y navegación de Administración')
h2('Ingresar con el rol correcto')
capture(19,'Ingreso administrativo: identificador y contraseña en el formulario propio del rol.')
steps(['En Iniciar sesión pulsá Acceso administrativo.','En el formulario separado escribí Identificador de administrador y Contraseña. Mostrar permite revisar la contraseña antes de enviarla.','Pulsá Ingresar a administración. La cabecera identifica la sesión como Administrador; el menú global agrega Administración.'])
outcome('Se abre el panel con Usuarios activo. Si no tenés una cuenta administradora, el acceso normal de estudiante no revela este panel.')
h2('Cómo se organiza el panel')
capture(20,'Administración: menú lateral con los siete módulos y cabecera del administrador.')
p('A la izquierda está el menú de módulos. Gestión contiene Usuarios, Diseños y Actividad; Juego contiene Experiencia de juego, Fuentes del mapa y Unidades y equivalencias; Sistema contiene Configuración. Al seleccionar un módulo cambian el título y el contenido central, sin salir del panel. El nombre o identificador de la cabecera permite comprobar con qué cuenta estás operando.')
table(['Rol','Pantallas y alcance'],[
('Jugador','Inicio, Niveles según avance, Aprendizaje, Reglas, Ranking, Mis diseños, perfil, Edición y Simulación de sus propios diseños.'),
('Administrador','Además de las vistas generales, abre Administración y sus siete módulos; puede gestionar cuentas, diseños, versiones de niveles y mantenimiento.')],[1.45,5.1])
p('La observación de esta guía fue de consulta: no se cambiaron roles, diseños, borradores, publicaciones, registros ni configuración. Los pasos de escritura siguientes describen controles implementados y el resultado previsto, no una operación ejecutada en la cuenta real.')

h1('10 Usuarios')
h2('Buscar y reconocer una cuenta')
capture(21,'Usuarios: buscador, filtro de rol y tabla de resultados, con datos personales ocultados.')
p('La pantalla Gestión de usuarios presenta una tabla con Nombre completo, Correo electrónico, Rol y Acciones. Buscar usuario acepta nombre o correo y Rol permite elegir Todos los roles, Administrador o Jugador. El filtro sólo reduce las filas visibles; no altera las cuentas. Cuando no hay coincidencias se indica «No hay usuarios que coincidan con el filtro».')
steps(['Abrí Administración → Usuarios.','Escribí parte del nombre o correo en Buscar usuario.','Si necesitás limitar por permiso, elegí el valor de Rol.','Comprobá el nombre y correo de la fila antes de usar una acción.'])
outcome('La tabla muestra solamente los usuarios que cumplen ambos filtros.')
h2('Cambiar un rol')
p('Cada fila tiene su propio selector de rol y botón Guardar. Seleccionar Administrador o Jugador prepara el valor de esa fila; Guardar envía el cambio. La posibilidad de acceder a Administración depende del rol guardado, no de la pantalla de ingreso elegida. Revisá la identidad de la fila porque un cambio de privilegio afecta el acceso posterior de esa persona.')
steps(['Localizá la cuenta con los filtros.','En su fila elegí Administrador o Jugador.','Pulsá Guardar en esa misma fila y leé el mensaje de confirmación o error.'])
outcome('La fila refleja el rol guardado si la solicitud fue aceptada.')
h2('Editar datos y eliminar una cuenta')
capture(22,'Editar usuario: campos del diálogo y botones Cerrar y Guardar, con datos ocultados.')
p('Editar abre el diálogo «Editar usuario». Allí están Nombre, Apellido, Correo electrónico y Nueva contraseña. Para una cuenta administradora también aparece Identificador de administrador. La clave puede quedar vacía para conservar la actual; si se cambia, el formulario exige al menos seis caracteres, una mayúscula y un carácter especial. Guardar confirma los campos; Cerrar abandona el diálogo.')
p('Eliminar solicita confirmación y da de baja la cuenta de forma definitiva. No lo uses para corregir un nombre o contraseña. En este recorrido sólo se abrieron las pantallas de consulta.')

h1('11 Diseños y actividad')
h2('Buscar diseños de jugadores')
capture(23,'Diseños: buscador, grupo por jugador y tarjetas, con nombres y correos ocultados.')
p('La pantalla Diseños agrupa las tarjetas por propietario. Cada grupo muestra nombre, correo y cantidad de diseños. Buscar diseños por jugador filtra por esos datos. Una tarjeta identifica Diseño y, si corresponde, la actividad asociada. Abrir muestra el detalle; Eliminar borra el diseño tras una confirmación.')
steps(['Abrí Administración → Diseños.','Filtrá por nombre o correo del jugador.','En la tarjeta adecuada pulsá Abrir.','Leé el encabezado del detalle para confirmar propietario e identificador. Revisá las secciones Líneas de metro, Estaciones, Paradas por línea, Tramos entre estaciones y Unidades de metro.','Pulsá Cerrar para volver al listado.'])
outcome('El detalle muestra los elementos de la red. Si una sección no tiene registros, aparece «No hay elementos registrados».')
h2('Acciones disponibles sobre el detalle')
capture(24,'Detalle de diseño: secciones de líneas, estaciones, paradas, tramos y unidades.')
p('En cada sección del detalle hay Agregar, Editar y Eliminar según el tipo. Agregar línea solicita nombre; Agregar estación solicita nombre y coordenadas; Agregar parada a línea relaciona línea y estación; Agregar tramo conecta dos estaciones; Agregar unidad de metro asocia unidad y línea. Editar abre preguntas sobre los campos correspondientes. Aplicá estos controles sólo tras comprobar el diseño, porque modifican la red del jugador.')
p('Eliminar diseño es más amplio que quitar un elemento: la confirmación advierte que puede borrar elementos, intento, puntaje y resultados, incluso en un nivel completado, y modificar progreso y ranking del propietario. Esa acción no se puede deshacer.')
h2('Consultar el registro de Actividad')
capture(25,'Actividad: columnas Fecha, Administrador, Acción y Detalle, sin datos sensibles.')
p('Actividad muestra Fecha, Administrador, Acción y Detalle. Es un registro de acciones administrativas, no una tabla de partidas de estudiantes. Cada fila ofrece Eliminar y el encabezado ofrece Eliminar todo. Ambos piden confirmación y borran entradas del registro; para una consulta basta con leer las filas. Si la tabla está vacía se muestra un estado vacío.')
steps(['Elegí Administración → Actividad.','Leé fecha, responsable, acción y detalle de la fila que buscás.','Volvé al módulo que necesitás desde el menú lateral.'])
outcome('La revisión del registro no altera ninguna entrada.')

h1('12 Experiencia de juego')
p('Este módulo administra la próxima versión de cada nivel. Su lista inicial permite elegir un nivel; Editar abre un borrador con siete pestañas. Cambiar campos en el editor todavía no publica nada. Guardar conserva el borrador; Previsualizar comprueba el contenido guardado; Publicar requiere una validación viable y confirmación editorial. La publicación se aplica a partidas futuras según la versión del nivel.')
h2('Elegir nivel y entender las acciones generales')
capture(26,'Experiencia de juego: lista de niveles y cabecera del editor con Guardar, Previsualizar y Publicar.')
steps(['En Administración → Experiencia de juego localizá el nivel 01 a 10.','Pulsá Editar. Comprobá el número y nombre en la cabecera.','Recorré las pestañas antes de modificar: Desafío, Metas y POI, Reglas, Red de referencia, Tarjetas, Vista previa e Historial.','Usá Salir para volver a la lista. Si hiciste cambios, Guardar crea una revisión del borrador.'])
outcome('Se conserva la separación entre borrador y versión publicada; la cabecera informa que los cambios siguen sin publicar.')
h2('Desafío')
capture(27,'Desafío: nombre, dificultad, relato, objetivo e instrucciones.')
p('Define la tarjeta que presenta el nivel: Nombre, Dificultad, Relato, Objetivo e Instrucciones. Relato sitúa la historia; Objetivo resume qué debe lograr la red; Instrucciones explican qué construir. Revisá estos cinco campos juntos para evitar una consigna que no coincida con las metas configuradas.')
h2('Metas y POI')
capture(28,'Metas y POI: una meta y un lugar objetivo con sus controles.')
p('La pestaña reúne las metas que el jugador verá y los lugares del mapa que debe cubrir. Una meta identifica la acción comprobada, cantidad y texto visible. Los POI agregan referencias territoriales y pueden exigir cercanía de una estación y conexión al recorrido. La vista como jugador ayuda a leer la consigna desde su perspectiva. Si ajustás un POI o su radio, comprobá después el resultado de la validación.')
h2('Reglas')
capture(29,'Reglas administrativas: cantidades, condiciones y selector de objetivo específico.')
p('Cantidades de la red fija mínimos como estaciones, líneas, conexiones, unidades o transbordos. Condiciones de funcionamiento y Objetivos específicos describen requisitos que la red o la simulación deben cumplir. El selector Regla del nivel permite revisar una condición concreta o una regla admitida para agregar. La pestaña también muestra las herramientas habilitadas; una consigna no debería pedir una acción que el jugador tiene desactivada.')
h2('Red de referencia')
capture(30,'Red de referencia: mapa privado y selector de elementos.')
p('Es una red privada de comprobación del nivel. Contiene mapa, estaciones, líneas, conexiones, tramos, unidades y ejecuciones de prueba. Elegí «Elementos de referencia», después «Elemento de referencia» para inspeccionar sus datos. Agregar incorpora un elemento al borrador; Aplicar actualiza el mapa con los campos editados; Quitar lo retira. Esta red sirve para comprobar que la consigna pueda resolverse, no aparece como diseño del jugador.')
h2('Tarjetas')
capture(31,'Tarjetas: selector, texto, fuente e imagen de una de las siete piezas.')
p('Cada nivel reúne siete tarjetas educativas. El selector permite pasar de una a otra. En cada ficha revisá título, texto, aprendizaje, fuente, URL de fuente e imagen. Guardar rechaza una URL de fuente que no sea HTTP o HTTPS válida. Antes de publicar hay una casilla editorial: «Revisé las afirmaciones, fuentes e imágenes de las siete tarjetas». Esa confirmación no sustituye la lectura efectiva de todas ellas.')
h2('Vista previa')
capture(32,'Vista previa: consigna, resultado de validación o tarjeta.')
steps(['Guardá el borrador antes de previsualizar. Si hay cambios sin guardar, la pantalla te pide guardarlos.','Pulsá Previsualizar.','En Vista previa alterná Consigna, Resultado de la validación y las siete tarjetas.','Leé las condiciones Cumplido o Pendiente y el mensaje diagnóstico. Corregí el borrador si la red de referencia no demuestra que se pueda completar el nivel.'])
outcome('Publicar sólo se habilita para una vista previa viable, sin cambios posteriores sin guardar. Previsualizar comprueba; no publica.')
h2('Historial y publicación')
capture(33,'Historial: versiones del nivel y navegación entre páginas.')
p('Historial lista versiones en páginas y permite consultar publicaciones anteriores. El control de limpieza elimina versiones antiguas sin uso sólo tras una confirmación; conserva la versión vigente, el borrador y las vinculadas a partidas. No es necesario para publicar.')
p('Para publicar, revisá las siete tarjetas, marcá la confirmación editorial, guardá, previsualizá y asegurate de que el diagnóstico sea viable. Entonces pulsá Publicar. Si cambiaste algo después de previsualizar, repetí Guardar y Previsualizar. Publicar crea una versión nueva del contenido completo; no se hizo durante la preparación de esta guía.')

h1('13 Fuentes, unidades y configuración')
h2('Fuentes del mapa')
capture(34,'Fuentes del mapa: procedencia y fecha de las referencias territoriales.')
p('La pantalla identifica procedencia y antigüedad de calles y avenidas, tráfico histórico, densidad de población y desplazamientos; los enlaces llevan a la fuente correspondiente. Sirve para interpretar las capas de Edición y Modo Libre. El tráfico mostrado es histórico, no una medición en vivo. La consulta de fuentes no cambia el mapa.')
h2('Unidades y equivalencias')
capture(35,'Unidades y equivalencias: valores UT y UV y ejemplo de 12 km.')
p('La tabla indica 1 UT = 1 hora y 1 UV = 1 km por UT, equivalente a 1 km/h en la escala didáctica. Incluye un recorrido de 12 km para comparar combinaciones de velocidad y duración y aclara que la cobertura se mide a 500 m en línea recta desde una estación. No es un formulario para modificar valores de una partida.')
p('Hay una contradicción visible: el texto estático de esta pantalla afirma que UV y UT son independientes en niveles, mientras la implementación vigente de la campaña los vincula por línea. La regla operativa en la simulación actual es UT = distancia ÷ UV, o UV = distancia ÷ UT, según el valor que se aplique. Se conserva esta discrepancia documentada hasta que el proyecto la resuelva.')
h2('Configuración y mantenimiento')
capture(36,'Configuración: estado de Modo de mantenimiento y botón Guardar.')
p('Configuración muestra Modo de mantenimiento y, si el servidor proporciona otras claves visibles, «Otros parámetros». El selector de mantenimiento ofrece ACTIVADO y DESACTIVADO. Al activarlo, los jugadores no pueden ingresar y los administradores mantienen acceso. Elegir un valor no lo aplica hasta pulsar Guardar. Si el guardado falla, el mensaje de estado explica el error. En la sesión consultada estaba desactivado y no se cambió.')
steps(['Abrí Administración → Configuración.','Leé el estado actual antes de editar.','Si la operación está autorizada para tu entorno, elegí el nuevo estado y pulsá Guardar.','Comprobá el mensaje y volvé a leer el estado mostrado.'])
outcome('El estado guardado se refleja en la pantalla y determina el acceso de jugadores.')

heading10=d.add_heading('14 Problemas habituales',1)
heading10.paragraph_format.page_break_before=True
table(['Situación','Qué comprobar'],[
('No abre la aplicación local','Verificá con quien administra la instalación que frontend y backend estén iniciados; la aplicación usa normalmente 5173 y 8080.'),
('No puedo iniciar sesión','Revisá correo, contraseña y rol de acceso. Usá recuperación si olvidaste la clave; no compartas contraseñas.'),
('Un nivel figura bloqueado','Comprobá el progreso en Niveles y terminá el anterior con Finalizar red.'),
('Simular está deshabilitado','Completá una red continua, asigná el metro requerido y revisá las metas y mensajes del editor.'),
('Un POI sigue pendiente','Ubicá la estación dentro del radio indicado y verificá que pertenezca al recorrido o línea exigidos.'),
('El transbordo no cuenta','Las dos líneas deben compartir una estación; el cruce visual de tramos no alcanza.'),
('No se aprueba tras simular','Volvé a Edición y elegí Finalizar red. Revisá los objetivos pendientes y la ejecución registrada.'),
('Una capa no carga','Reintentá con el control de la capa. La construcción puede continuar mientras falta esa referencia.'),
('No veo Extras','Primero ingresá a Modo Libre; después abrí Aprendizaje → Modo Libre.')],[2.0,4.55])

h1('15 Glosario')
table(['Término','Significado en METRONET'],[
('Estación','Punto de parada que podés ubicar en el mapa.'),('Línea','Recorrido que agrupa estaciones y conexiones.'),('Conexión o tramo','Enlace entre dos estaciones de una línea.'),('Metro','Unidad que circula por una línea diseñada.'),('Transbordo','Estación compartida por dos líneas.'),('POI','Punto de interés usado como referencia u objetivo territorial.'),('UV','Unidad de velocidad; 1 UV equivale a 1 km por UT.'),('UT','Unidad de tiempo simulado; 1 UT equivale a una hora.'),('Cobertura','Radio de proximidad representado alrededor de una estación.'),('Intento','Partida de un nivel con diseño, simulaciones y posible resultado.'),('Extras','Siete tarjetas didácticas disponibles al entrar a Modo Libre.')],[1.65,4.9])

d.core_properties.title='Manual de usuario METRONET'
d.core_properties.subject='Guía para estudiantes de logística y administración'
d.core_properties.author='METRONET'
d.save(OUT)
print(OUT)

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
OUT = BASE / 'Manual integral METRONET 2026-10-18 rev2.docx'
M = BASE / 'materiales'

d = Document()
s = d.sections[0]
s.top_margin = Inches(.74); s.bottom_margin = Inches(.7)
s.left_margin = Inches(.83); s.right_margin = Inches(.83)
s.header_distance = Inches(.3); s.footer_distance = Inches(.35)

styles = d.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor(28, 34, 43)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.15
for name, size, before, after in [('Title', 24, 0, 12), ('Heading 1', 15, 16, 8), ('Heading 2', 11.5, 11, 5), ('Heading 3', 10.5, 8, 4)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name in ('Title','Heading 1') else 'Aptos'
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
styles['Caption'].font.name = 'Aptos'; styles['Caption'].font.size = Pt(8.5)
styles['Caption'].font.color.rgb = RGBColor(70, 78, 89)
styles['Caption'].paragraph_format.space_after = Pt(9)
title_style_ppr = styles['Title']._element.get_or_add_pPr()
border = title_style_ppr.find(qn('w:pBdr'))
if border is not None: title_style_ppr.remove(border)

footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('METRONET  ·  Manual de usuario  ·  18.10.2026  ·  ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def title(t): d.add_paragraph(t, 'Title')
def h1(t): d.add_heading(t, 1)
def h2(t): d.add_heading(t, 2)
def p(t): d.add_paragraph(t)
def steps(items):
    for i, item in enumerate(items,1):
        q=d.add_paragraph(style='Normal'); q.paragraph_format.left_indent=Inches(.22); q.paragraph_format.first_line_indent=Inches(-.22)
        r=q.add_run(f'{i}. '); r.bold=True; q.add_run(item)
def bullet(items):
    for item in items:
        q=d.add_paragraph(item, style='List Bullet'); q.paragraph_format.space_after=Pt(3)
def outcome(t):
    q=d.add_paragraph(); q.add_run('Resultado esperado. ').bold=True; q.add_run(t)
def fig(filename, caption, width=4.65):
    img=M/filename
    if img.exists():
        q=d.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.keep_with_next=True
        shape=q.add_run().add_picture(str(img), width=Inches(width))
        shape._inline.docPr.set('descr', caption)
        c=d.add_paragraph(caption, 'Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style='Table Grid'
    for i, x in enumerate(headers): t.rows[0].cells[i].text=x
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row): c[i].text=str(x)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    trPr=t.rows[0]._tr.get_or_add_trPr(); header=OxmlElement('w:tblHeader'); header.set(qn('w:val'),'true'); trPr.append(header)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.font.size=Pt(8.5)
            tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'), '1E3656' if ri==0 else ('F1F5F8' if ri%2==0 else 'FFFFFF')); tcPr.append(sh)
            if ri==0:
                for para in cell.paragraphs:
                    for run in para.runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(1)

title('Manual de usuario METRONET')
p('Guía de la versión local verificada el 18 de octubre de 2026. Destinada a estudiantes de logística que diseñan una red, ponen a prueba sus recorridos y comparan decisiones de operación. La administración se explica al final para quienes tengan ese rol.')
fig('logo.png','Figura 1. Logotipo original de METRONET incluido en la aplicación.',3.1)
p('Versión del repositorio: b5168659bebe276fee89f37963b97daf22992503. Edición del manual: 1.0 · 18 de octubre de 2026.')
p('Estado de esta guía: IMPLEMENTADO describe la interfaz y el código de la versión local. Se observaron en vivo acceso, navegación, Edición y administración; la red vacía impidió entrar a Simulación, cuya descripción se verificó en el código. La única imagen incorporada es el logotipo actual; las capturas de pantalla siguen pendientes.')

h1('Índice')
for line in ['PARTE A · USUARIO','1 Registro, acceso, recuperación y perfil','2 Campaña y niveles','3 Tutoriales y Extras','4 Edición de la red','5 Mapa, capas, objetivos y POI','6 Guardado, simulación y UV y UT','7 Resultados, diseños, aprendizaje y reglas','8 Ranking, trofeos y Modo Libre','PARTE B · ADMINISTRADOR','9 Acceso y siete módulos administrativos','10 Problemas habituales','11 Glosario']:
    p(line)

h1('PARTE A · USUARIO')
p('Los capítulos 1 a 8 describen la experiencia de una cuenta jugadora. El administrador también puede consultar varias de estas pantallas, pero su cuenta tiene accesos adicionales y no participa de la clasificación como jugador.')
h1('1 Registro, acceso, recuperación y perfil')
h2('Crear una cuenta de estudiante')
steps(['Desde Iniciar sesión elegí Crear cuenta.','Completá Nombre, Apellido, Correo electrónico y Contraseña. La clave necesita al menos seis caracteres, una mayúscula y un carácter especial; Mostrar permite verla mientras escribís.','Leé Leer el uso de mis datos y aceptá la casilla de privacidad si estás de acuerdo.','Pulsá Crear cuenta. Después ingresá por el acceso de estudiante.'])
outcome('Se crea una cuenta de jugador cuando los datos son válidos y el correo está disponible. Crear cuenta no habilita funciones de administrador.')
h2('Entrar como estudiante')
steps(['Abrí la dirección local que usa el proyecto, normalmente http://127.0.0.1:5173/login.html, con frontend y backend ya iniciados.','Escribí tu correo electrónico y contraseña. El botón Mostrar permite comprobar lo escrito.','Elegí Iniciar sesión. La aplicación conduce a Inicio o al destino protegido desde el que llegaste.'])
outcome('Se ve el inicio con acceso al nivel actual y al Modo Libre si está habilitado.')
h2('Recuperar el acceso')
steps(['En Iniciar sesión elegí ¿Olvidaste tu contraseña?','En Recuperar acceso escribí el correo de tu cuenta y pulsá Enviar código. La pantalla evita revelar si ese correo existe.','Si recibís el mensaje, abrí Verificar código, ingresá sus seis dígitos dentro del plazo indicado y pulsá Verificar código. Reenviar código se habilita tras la espera visible.','En Nueva contraseña escribí y repetí una clave válida; elegí Cambiar contraseña. Volvé al ingreso.'])
outcome('Tras una verificación válida, la contraseña nueva reemplaza a la anterior. Las dos últimas pantallas sólo aparecen si se inició y verificó una solicitud real; este manual no envió correos ni cambió credenciales.')
h2('Orientarse en la aplicación')
p('Para el estudiante, la navegación ofrece Inicio, Niveles, Aprendizaje, Reglas, Ranking y Mis diseños. El administrador ve además Administración. El menú del nombre abre Mi perfil y Cerrar sesión; el control Música cambia el sonido. Durante un nivel se indica la secuencia Nivel → Edición → Simulación → Resultados. Volvé a Inicio para retomar el recorrido que corresponda a tu cuenta.')
p('En Inicio, la tarjeta principal identifica el nivel para continuar y muestra el estado de la campaña. La segunda tarjeta permite entrar a Modo Libre cuando se desbloquea. Si necesitás consultar o cambiar tus datos, abrí Mi perfil desde el menú de usuario; allí se separan datos personales, correo y contraseña.')
h2('Mi perfil')
p('Mi perfil muestra rol y fecha de ingreso. Datos personales permite actualizar nombre y apellido; Correo electrónico muestra el correo actual y recibe uno nuevo; Contraseña solicita clave actual, nueva y confirmación, con botones Mostrar y los requisitos visibles. Cada bloque tiene su propio Guardar. Volver a Inicio sale del perfil. Cambiar estos datos modifica la cuenta, por eso el recorrido de verificación sólo observó los formularios.')

h1('2 Campaña y niveles')
h2('Elegir y continuar un nivel')
steps(['Entrá en Niveles desde el menú.','Seleccioná un número del 01 al 10. La tarjeta muestra nombre, estado, relato y desafío.','Leé el objetivo antes de iniciar o retomar. Elegí el botón de la tarjeta cuando el nivel esté desbloqueado.','La preparación muestra el relato y una animación de viaje. Podés esperar o pulsar Iniciar nivel. Cuando aparece una tarjeta educativa, leé la idea clave y pulsá Continuar para abrir Edición.'])
outcome('Se abre una red del intento actual y aparecen las metas que debés cumplir. Un nivel bloqueado no permite iniciarlo.')
p('La pantalla Niveles tiene un selector numerado y una tarjeta de detalle. En el selector leés estados como disponible, en curso, completado o bloqueado. En el detalle aparece el botón Iniciar, Continuar, Repetir o Bloqueado según el estado. La preparación muestra el relato y las metas antes de entrar al plano.')
p('La campaña progresa de forma secuencial. Los diez niveles enseñan de un primer enlace entre estaciones a redes de dos líneas, cobertura de lugares, transbordos y comparación de simulaciones. Completar una simulación no aprueba por sí solo el nivel: hay que volver a Edición y usar Finalizar red.')
p('En la cuenta administradora observada, el catálogo ofrece acceso a los diez niveles y Modo Libre. Eso es una vista del rol de administración: para un jugador la disponibilidad depende de su propio avance y no debe interpretarse como un desbloqueo general.')
table(['Niveles','Qué practicás'],[
('1 a 3','Primer enlace, extensión, velocidad y duración simulada.'),
('4 a 6','Varias unidades, cobertura de POI y continuidad de la línea.'),
('7 a 8','Dos líneas, transbordo y objetivo territorial.'),
('9 a 10','Comparaciones de ejecuciones y desafío integrador.')],[1.2,5.35])
p('Reiniciar recorrido aparece sólo si hay avance en la campaña. Abre una confirmación; usalo deliberadamente porque cambia el recorrido actual, aunque el sistema conserva el histórico según sus verificaciones. Este manual no ejecutó esa acción.')

h1('3 Tutoriales y Extras')
h2('Seguir o volver a abrir una ayuda')
steps(['Al abrir Edición por primera vez, elegí Iniciar tutorial si querés la guía paso a paso, o Continuar para construir directamente.','Avanzá leyendo la instrucción y realizá la acción indicada en la interfaz. Los pasos prácticos comprueban acciones reales, no sólo pulsar Siguiente.','Para volver a verla, abrí el botón Tutorial de la esquina superior derecha y elegí Iniciar tutorial. La Guía rápida resume Estación, Conexión, Metro, Guardar, Simular y Finalizar.'])
outcome('El panel de ayuda cambia de paso cuando se cumple la acción esperada. Las ayudas extensas se presentan al introducir herramientas nuevas y quedan disponibles manualmente.')
p('La campaña presenta la construcción en el nivel 1, POI en el 5 y transbordo en el 7. La simulación introduce controles en los niveles correspondientes. Las guías pueden variar según la versión publicada del intento.')
h2('Extras de Modo Libre')
p('Al entrar por primera vez a Modo Libre, la aplicación puede ofrecer siete tarjetas Extras. Explican calles, población, tráfico, desplazamientos, horarios, cobertura y relación entre diseño y simulación. Podés volver a consultarlas en Aprendizaje → Modo Libre. Su lectura es material didáctico; no modifica la velocidad, la duración ni la red.')
p('Cada Extra muestra un título, qué ofrece la herramienta, cómo ayuda a la red, una práctica y una advertencia de interpretación. Podés avanzar por las siete tarjetas, salir y volver desde la biblioteca.')

h1('4 Edición de la red')
h2('Crear un recorrido básico')
steps(['Desde un nivel desbloqueado, entrá en Edición y leé Metas del nivel.','Elegí Estación en Construcción y hacé clic en un lugar válido del mapa. Repetí para ubicar otra estación.','Elegí Línea y creá la línea que vas a usar.','Elegí Conexión, seleccioná la línea activa y uní dos estaciones. Para extenderla, conectá estaciones adicionales de la misma línea.','Elegí Metros y hacé clic sobre un tramo inequívoco para asignar una unidad. Si hay líneas superpuestas, elegí la que corresponde.','Pulsá Guardar y comprobá el aviso Diseño guardado.'])
outcome('El panel presenta estaciones, líneas, tramos y metros del diseño; las metas se actualizan de acuerdo con la red confirmada.')
p('La herramienta Seleccionar permite inspeccionar una estación, línea, conexión o metro. Desde el elemento seleccionado podés editar sus datos; Mover y Eliminar se habilitan cuando corresponde. Cancelar operación devuelve al modo normal. Una eliminación abre un diálogo porque puede quitar conexiones u otras partes de la red.')
p('En Edición distinguí tres zonas: el mapa a la izquierda, Metas del nivel a la derecha (y Lugares a conectar cuando la consigna usa POI), y Herramientas debajo de las metas. Construcción reúne Estación, Línea, Conexión y Metros; Edición reúne Seleccionar, Mover, Eliminar y Cancelar. Acciones reúne Guardar, Simular diseño y, para niveles, Finalizar red. El aviso de estado comunica confirmaciones y errores.')
p('En la visita al nivel 1, las cinco metas visibles fueron: ubicar dos estaciones, crear una línea, crear una conexión, asignar una unidad de metro y simular la red con sus UV/UT actuales. Sin ningún elemento creado, Mover, Eliminar y Cancelar estaban deshabilitados; Seleccionar era la herramienta activa. El mapa ofrecía zoom, Ajustar red y tres controles de orientación. El botón Aprender y el botón Tutorial estaban en la cabecera.')
p('Un transbordo requiere que dos líneas compartan una estación real. Dibujar dos vías que se cruzan sin esa estación no cuenta. En los niveles avanzados se exige al menos un metro por línea.')

h1('5 Mapa, capas, objetivos y POI')
h2('Leer el mapa')
p('El mapa de Montevideo muestra barrios, estaciones, líneas y referencias. Los controles + y − cambian el zoom; restaurar vista vuelve al encuadre de referencia. También hay mandos de orientación. La barra geográfica abre referencias y opciones del mapa. Al seleccionar una zona o un elemento, la interfaz muestra datos relacionados.')
h2('Usar POI y comprobar una meta')
steps(['Abrí Puntos de interés desde el control de referencias del mapa.','Usá Buscar punto de interés para localizar el lugar solicitado. Podés activar categorías de marcadores.','Seleccioná el POI para centrar la vista y observar su referencia.','Ubicá una estación dentro del radio que muestra la meta; conectala por la línea requerida.','Revisá Metas del nivel y Lugares a conectar. Confirmá que el estado cambie de Pendiente a Cubierto o Completado.'])
outcome('El objetivo se marca como cubierto cuando la geometría y las conexiones cumplen la consigna; estar visualmente cerca de la estrella no basta.')
p('En los niveles del recorrido, Palacio Legislativo y Torre de las Comunicaciones aparecen como referencias de cobertura. Sus radios de evaluación son reglas de juego y no equivalen directamente a kilómetros. Leé siempre la consigna publicada de tu intento.')
h2('Capas territoriales')
p('Modo Libre habilita capas de calles, tráfico histórico, población y desplazamientos, además de cobertura y controles horarios. Activá una capa desde Análisis del territorio, ajustá opacidad cuando exista ese control y compará el mismo trazado con distintas vistas. En Simulación libre hay un panel Capas del mapa. La cobertura dibuja 500 m en línea recta alrededor de cada estación: no calcula accesibilidad peatonal ni demanda real. Tráfico y desplazamientos son referencias históricas; no representan datos en vivo.')
p('El selector horario cambia la hora de referencia de tráfico y desplazamientos. La población permanece fija. Si una capa falla, el panel muestra un aviso y permite reintentar; podés seguir trabajando con la red.')

h1('6 Guardado, simulación y UV y UT')
h2('De Edición a Simulación')
steps(['En Edición, pulsá Guardar para conservar la red y revisar las metas.','Cuando la red sea consistente y tenga el metro requerido, elegí Simular diseño. La aplicación comprueba y guarda la preparación antes de abrir Simulación.','En Simulación, observá el mapa, la ficha de unidad y los controles. Elegí Iniciar.','Usá Pausar, Reanudar, Detener o Reiniciar si necesitás revisar la ejecución. Al finalizar, consultá el resultado y regresá a Edición para Finalizar red.'])
outcome('Una ejecución terminada queda registrada. Consultar sus resultados no aprueba la campaña; Finalizar red realiza esa comprobación.')
p('Comprobación en vivo: al pulsar Simular diseño con la red vacía del nuevo intento, apareció «Prepará la red para simular» con la indicación de completar objetivos y asignar un metro. No se abrió Simulación ni se registró una ejecución. La descripción de la pantalla y controles que sigue corresponde a la implementación inspeccionada en el código, no a una visita visual de esa pantalla.')
p('La pantalla Simulación tiene un mapa y un panel de instrumentos. En Acciones están Iniciar, Pausar, Reanudar, Detener y Reiniciar. Debajo se ven Estado, Tiempo simulado y el porcentaje de Ejecución. El panel de parámetros contiene Duración · UT y, cuando corresponde, la velocidad UV, el selector de línea o unidad y el ritmo visual. Ampliar mapa oculta temporalmente parte de los instrumentos para ver mejor el recorrido.')
h2('Ajustar una línea')
p('UV expresa velocidad: 1 UV = 1 km por UT. UT expresa tiempo simulado: 1 UT = 1 hora del escenario. En el recorrido vigente, cada línea comparte UV y UT entre sus metros. Al asignar el primer metro, comienza con 1 UV. Si elegís UV, la duración de la línea se calcula como distancia ÷ UV; si elegís UT, la velocidad se calcula como distancia ÷ UT. Otra línea conserva sus propios valores.')
steps(['Seleccioná una línea o unidad según el control disponible en el panel. Comprobá qué línea indica la ficha antes de editar.','Cambiá Velocidad · UV o Duración · UT a un valor mayor que cero. Usá los botones + y − o escribí el número.','Elegí el botón de guardar velocidad o Aplicar UT. Observá el otro valor recalculado para la distancia de esa línea.','Volvé a ejecutar para comparar el resultado. Si el nivel pide una comparación, conservá las demás variables que indique su consigna.'])
outcome('El panel muestra UV y UT de la línea seleccionada y la ejecución nueva refleja el ajuste guardado.')
p('Ritmo visual × controla la reproducción en pantalla; no equivale a UV. Una línea más corta puede terminar antes que otra y quedar en su estación terminal. Los valores mostrados pueden redondearse a dos decimales; la aplicación guarda más precisión internamente.')
table(['Acción','Efecto'],[('Guardar','Conserva la red y actualiza metas.'),('Simular','Comprueba, ejecuta y registra una prueba de la red.'),('Finalizar red','Evalúa la consigna y, si corresponde, aprueba el nivel.')],[1.45,5.1])

h1('7 Resultados, diseños, aprendizaje y reglas')
h2('Interpretar el resultado')
p('Simulación muestra estado, tiempo simulado, porcentaje de ejecución, unidad seguida y mensajes de la red. Tras una ejecución se puede consultar el resultado guardado. Al finalizar un nivel aprobado aparece una vista con puntos obtenidos, posible récord personal y desglose del cálculo. Abrí cada ejecución del desglose para conocer descuentos o condiciones pendientes; elegí Continuar para seguir.')
h2('Volver a un diseño')
steps(['Entrá en Mis diseños.','Elegí Diseño de los niveles o Modo Libre. Las pestañas separan ambas clases de red.','En una tarjeta propia, elegí Abrir diseño para editar o Simular diseño si está preparado. Usá Anterior y Siguiente si la lista tiene varias páginas.'])
outcome('Se abre el diseño elegido con su contexto. La opción Simular puede estar deshabilitada si la red aún no está preparada.')
p('Eliminar diseño pide confirmación y no se puede deshacer. En un diseño de nivel también puede quitar partida, puntaje y resultados, con efecto sobre progreso y ranking. No lo uses como forma de volver a editar.')
h2('Consultar Aprendizaje y Reglas')
p('Aprendizaje permite elegir un nivel y abrir sus siete tarjetas publicadas cuando está desbloqueado. Cada tarjeta incluye imagen, explicación, idea clave y fuente. La pestaña Modo Libre reúne las siete tarjetas Extras luego de ingresar a ese modo.')
p('Reglas del juego se organiza en Puntuación, Cómo aprobar, Acciones, Velocidad y tiempo, y Ranking. Consultá Cómo aprobar antes de Finalizar red; en Velocidad y tiempo se explican UV y UT. La pestaña Puntuación detalla descuentos y cálculo. Si una regla de una publicación histórica difiere, la vista de resultado puede mostrar la regla original de ese intento.')
p('En Mis diseños, cada tarjeta incluye nombre, estado y acciones para abrir o simular cuando estén disponibles. Los contadores de las pestañas indican cuántos diseños hay en cada grupo. En Aprendizaje, el selector 01 a 10 y Modo Libre cambia las tarjetas visibles. Las pantallas pueden indicar bloqueado o error de carga; Reintentar aparece cuando la consulta falla.')

h1('8 Ranking, trofeos y Modo Libre')
h2('Ranking y reconocimientos')
p('En Ranking podés alternar Clasificación y Puntos por nivel. Hay navegación entre páginas y secciones Trofeos y Tu desempeño. La mejor marca de un nivel cuenta para la clasificación según la regla visible en Reglas. Los trofeos aparecen cuando se cumplen las condiciones correspondientes; el panel muestra el estado alcanzado para tu cuenta.')
p('Clasificación ordena jugadores y sus puntos; Puntos por nivel desglosa los diez niveles en páginas. Trofeos permite seleccionar un reconocimiento para leer el requisito, incluso cuando todavía está oculto. Tu desempeño resume tus marcas. Si no hay puntuaciones, la pantalla muestra un estado vacío en lugar de una clasificación inventada.')
p('Los reconocimientos visibles son Primera estación, Camino a la cima, Copa de la red, Biblioteca completa, Corona de platino y Premio oculto. Seleccionar una ficha muestra su requisito. En la cuenta administradora observada figuran «En exhibición» y el aviso de que esa vista no atribuye premios; los logros reales corresponden a cuentas jugadoras.')
h2('Diseñar sin una consigna de campaña')
steps(['Completá los diez niveles del recorrido. El acceso a Modo Libre aparecerá disponible en Inicio, Niveles y Mis diseños.','Elegí Modo Libre y creá una red nueva. Si se ofrece la bienvenida de Extras, leé o recorré sus tarjetas.','Ubicá estaciones, construí líneas y conexiones, asigná metros y guardá.','Explorá las capas territoriales y la cobertura. Probá UV o UT por línea, simulá y compará tu decisión con el resultado.'])
outcome('La red queda en la pestaña Modo Libre de Mis diseños. Podés editarla y volver a simularla sin una meta de aprobación de nivel.')
p('Modo Libre sirve para ensayar decisiones de planificación. Las capas no calculan por sí solas demanda, capacidad ni viabilidad real.')

heading9=d.add_heading('PARTE B · ADMINISTRADOR',1)
heading9.paragraph_format.page_break_before=True
p('Esta parte describe únicamente los permisos y pantallas adicionales de la cuenta administradora. Las operaciones comunes de consulta se explican en la Parte A; los botones administrativos pueden cambiar cuentas, contenidos y acceso a la plataforma.')
h1('9 Acceso y siete módulos administrativos')
h2('Ingresar y reconocer el rol')
p('Desde Iniciar sesión elegí Acceso administrativo. La pantalla separada pide Identificador de administrador y Contraseña; Mostrar permite revisar la entrada. Elegí Ingresar a administración. Una cuenta de estudiante no obtiene el menú Administración. El administrador ve ese menú adicional y puede consultar las pantallas comunes. En Ranking ve el catálogo de trofeos «En exhibición», sin que se le adjudiquen premios.')
p('El menú lateral agrupa Gestión (Usuarios, Diseños, Actividad), Juego (Experiencia de juego, Fuentes del mapa, Unidades y equivalencias) y Sistema (Configuración). Los jugadores no ven estas opciones.')
table(['Módulo','Uso visible'],[
('Usuarios','Buscar por nombre o correo, filtrar por rol y abrir edición de datos o privilegios.'),
('Diseños','Buscar diseños por jugador y consultar su detalle.'),
('Actividad','Consultar el registro de acciones administrativas.'),
('Experiencia de juego','Elegir un nivel y editar el borrador de su próxima publicación; incluye metas y referencia.'),
('Fuentes del mapa','Consultar fuentes territoriales mostradas por el sistema.'),
('Unidades y equivalencias','Consultar la escala didáctica de UV, UT y cobertura.'),
('Configuración','Consultar y administrar opciones de sistema disponibles.')],[1.75,4.8])
h2('Usuarios, diseños y actividad')
p('Usuarios muestra nombre, correo, rol y acciones. Buscar usuario filtra por nombre o correo; Rol reduce la lista. El selector de cada fila permite preparar un rol distinto y Guardar lo aplica. Editar abre los datos de esa cuenta; Eliminar inicia la baja. Estas acciones cambian cuentas reales: confirmá la fila antes de usarlas. En la visita sólo se consultó la lista.')
p('Diseños presenta los diseños de jugadores y un buscador por jugador. Si no hay registros, aparece el estado vacío. Actividad presenta fecha, administrador, acción y detalle; incluye Eliminar por fila y Eliminar todo para el registro. Esos botones no son pasos de consulta y no se usaron.')
h2('Experiencia de juego')
steps(['Elegí Experiencia de juego y el botón Editar del nivel 01–10 que querés revisar.','En Desafío se muestran nombre, dificultad, relato, objetivo e instrucciones.','Metas y POI enumera metas, acción, cantidad, texto y lugares objetivo; Reglas expone mínimos, condiciones y herramientas habilitadas.','Red de referencia permite definir elementos sobre un mapa privado; Tarjetas reúne las siete piezas educativas con texto, fuente e imagen.','Vista previa sirve para revisar un borrador guardado; Historial registra versiones. Guardar conserva el borrador y Publicar versión afecta a futuras partidas.'])
outcome('Al consultar las pestañas se identifican los contenidos de la próxima publicación sin cambiar lo que ve el jugador. No se guardó ni publicó nada durante esta revisión.')
table(['Pestaña','Campos y controles visibles'],[
('Desafío','Nombre, dificultad, relato, objetivo e instrucciones de la tarjeta del nivel.'),
('Metas y POI','Lista de metas; acción comprobada, cantidad y texto; lugares objetivo y vista como jugador.'),
('Reglas','Mínimos de red, exigencia de simulación, herramientas habilitadas y reglas admitidas.'),
('Red de referencia','Mapa privado, acción activa, línea, elementos, coordenadas y controles Aplicar o Quitar.'),
('Tarjetas','Selector de las siete tarjetas, título, texto, aprendizaje, fuente, imagen y revisión.'),
('Vista previa','Resultado del borrador luego de guardarlo y usar Previsualizar.'),
('Historial','Versiones registradas y controles de consulta del historial.')],[1.55,5.0])
h2('Fuentes, unidades y configuración')
p('Fuentes del mapa identifica procedencia y antigüedad de calles y avenidas, tráfico histórico, densidad de población y desplazamientos, con enlaces a cada fuente. El tráfico es una referencia histórica, no información en vivo. Unidades y equivalencias ofrece ejemplos de UV, UT y cobertura. Configuración contiene Modo de mantenimiento: al activarlo se impide el ingreso de jugadores y se conserva el acceso administrador. Cambiar el selector y pulsar Guardar modifica el servicio; aquí sólo se comprobó el estado desactivado.')
p('Discrepancia a tener en cuenta: la tarjeta estática Unidades y equivalencias de Administración todavía afirma que UV y UT son independientes en niveles. La implementación vigente de la campaña y su verificación del 15 de octubre establecen el vínculo automático por línea. Para operar un nivel, seguí el comportamiento del panel de Simulación y la consigna de tu intento.')

heading10=d.add_heading('10 Problemas habituales',1)
heading10.paragraph_format.page_break_before=True
table(['Situación','Qué comprobar'],[
('No abre la aplicación local','Verificá con quien administra la instalación que frontend y backend estén iniciados; la aplicación usa normalmente 5173 y 8080.'),
('No puedo iniciar sesión','Revisá correo, contraseña y rol de acceso. Usá recuperación si olvidaste la clave; no compartas contraseñas.'),
('Un nivel figura bloqueado','Comprobá el progreso en Niveles y terminá el anterior con Finalizar red.'),
('Simular está deshabilitado','Completá una red continua, asigná el metro requerido y revisá las metas y mensajes del editor.'),
('Un POI sigue pendiente','Ubicá la estación dentro del radio indicado y verificá que pertenezca al recorrido o línea exigidos.'),
('El transbordo no cuenta','Las dos líneas deben compartir una estación; el cruce visual de tramos no alcanza.'),
('No se aprueba tras simular','Volvé a Edición y elegí Finalizar red. Revisá los objetivos pendientes y la ejecución registrada.'),
('Una capa no carga','Reintentá con el control de la capa. La construcción puede continuar mientras falta esa referencia.'),
('No veo Extras','Primero ingresá a Modo Libre; después abrí Aprendizaje → Modo Libre.')],[2.0,4.55])

h1('11 Glosario')
table(['Término','Significado en METRONET'],[
('Estación','Punto de parada que podés ubicar en el mapa.'),('Línea','Recorrido que agrupa estaciones y conexiones.'),('Conexión o tramo','Enlace entre dos estaciones de una línea.'),('Metro','Unidad que circula por una línea diseñada.'),('Transbordo','Estación compartida por dos líneas.'),('POI','Punto de interés usado como referencia u objetivo territorial.'),('UV','Unidad de velocidad; 1 UV equivale a 1 km por UT.'),('UT','Unidad de tiempo simulado; 1 UT equivale a una hora.'),('Cobertura','Radio de proximidad representado alrededor de una estación.'),('Intento','Partida de un nivel con diseño, simulaciones y posible resultado.'),('Extras','Siete tarjetas didácticas disponibles al entrar a Modo Libre.')],[1.65,4.9])

d.core_properties.title='Manual de usuario METRONET'
d.core_properties.subject='Guía para estudiantes de logística y administración'
d.core_properties.author='METRONET'
d.save(OUT)
print(OUT)

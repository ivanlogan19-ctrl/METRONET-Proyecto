from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
ORIGEN = BASE.parent / "manual-usuario-2026-10-18-revision3" / "Manual METRONET Usuario y Administrador 2026-10-18 revision 3.docx"
CAPTURAS = BASE / "materiales" / "capturas_originales"
SALIDA = BASE / "Manual integral METRONET Administrador y Usuario 2026-10-18.docx"

# Todas las capturas son originales aportadas por la persona usuaria. El recorte
# se aplica sólo a la presentación en Word; los PNG permanecen intactos.
# Cada entrada: hora del archivo, explicación visible, recorte lateral opcional.
FIGURAS = {
    1: [
        ("6.57.29", "Registro de una cuenta y aceptación del uso de datos", 0, 0),
        ("6.57.33", "Detalle de la información de privacidad antes del registro", 0, 0),
    ],
    2: [("6.57.18", "Inicio de sesión de una cuenta jugadora", 0, 0)],
    3: [("6.57.41", "Solicitud de un código para recuperar el acceso", 0, 0)],
    6: [
        ("6.58.11", "Inicio con acceso al recorrido y al Modo Libre", 0, 0),
        ("6.58.15", "Menú de cuenta con Mi perfil y Cerrar sesión", 0, 0),
    ],
    7: [("6.58.29", "Mi perfil con bloques independientes de datos, correo y contraseña", 0, 0)],
    8: [("6.58.35", "Selección de los diez niveles y resumen del desafío", 0, 0)],
    9: [
        ("7.00.44", "Preparación del nivel 1 antes de entrar a Edición", 0, 0),
        ("7.00.48", "Tarjeta de aprendizaje presentada antes del nivel", 0, 0),
    ],
    10: [
        ("7.00.53", "Oferta de tutorial al entrar al editor", 0, 0),
        ("7.00.57", "Primer paso del tutorial en el mapa", 0, 0),
        ("7.01.02", "Ayuda contextual durante la construcción", 0, 0),
    ],
    11: [
        ("7.01.46", "Presentación de las siete tarjetas Extras del Modo Libre", 0, 0),
        ("7.01.49", "Extra sobre calles y avenidas", 0, 0),
    ],
    12: [("7.01.14", "Editor con estaciones, línea, unidad y metas del nivel", 0, 0)],
    13: [("7.01.14", "Detalle del mapa y controles de orientación", 0, 0.28)],
    14: [("7.02.00", "Puntos de interés y referencias superpuestos al mapa", 0, 0)],
    16: [("7.01.17", "Confirmación visual de Diseño guardado", 0.28, 0)],
    17: [
        ("7.01.22", "Aviso sobre UV y UT al preparar la simulación", 0, 0),
        ("7.01.32", "Tarjeta de aprendizaje abierta en la simulación", 0, 0),
    ],
    19: [
        ("7.02.05", "Mis diseños con el diseño de Nivel 1", 0, 0),
        ("7.02.08", "Mis diseños con el diseño de Modo Libre", 0, 0),
    ],
    20: [("6.58.38", "Biblioteca de aprendizaje y siete tarjetas del nivel 1", 0, 0)],
    21: [
        ("6.58.42", "Reglas de puntuación y descuentos", 0, 0),
        ("6.58.45", "Condiciones y pasos para aprobar", 0, 0),
        ("6.58.48", "Diferencia entre Guardar, Simular y Finalizar red", 0, 0),
        ("6.58.54", "Velocidad, tiempo y equivalencias dentro de las reglas", 0, 0),
        ("6.58.59", "Regla de la mejor marca en el ranking", 0, 0),
    ],
    22: [("6.59.03", "Ranking visto desde una cuenta administradora, con trofeos en exhibición", 0, 0)],
    23: [("7.01.52", "Editor de Modo Libre al iniciar una red propia", 0, 0)],
    24: [
        ("6.57.24", "Acceso separado para cuentas administradoras", 0, 0),
        ("6.57.55", "Pantalla de transición de METRONET al ingresar", 0, 0),
    ],
    25: [("6.59.18", "Navegación lateral de Gestión, Juego y Sistema", 0, 0)],
    26: [("6.59.10", "Búsqueda de usuarios, filtro de rol y acciones por fila", 0, 0)],
    28: [("7.02.19", "Supervisión de diseños agrupados por usuario", 0, 0)],
    31: [("6.59.18", "Lista de los diez niveles para edición administrativa", 0.16, 0)],
    32: [("6.59.21", "Editor del desafío con nombre, dificultad y consigna", 0.16, 0)],
    33: [("6.59.24", "Editor de metas y lugares objetivo", 0.16, 0)],
    34: [("6.59.27", "Cantidades, herramientas y condiciones configurables", 0.16, 0)],
    35: [("6.59.30", "Red de referencia y propiedades del elemento seleccionado", 0.16, 0)],
    36: [("6.59.33", "Tarjetas educativas con texto, imagen y fuente", 0.16, 0)],
    37: [("6.59.46", "Vista previa que solicita guardar el borrador", 0.16, 0)],
    38: [("6.59.51", "Historial de versiones del nivel", 0.16, 0)],
    39: [("7.00.19", "Fuentes de calles, tráfico, población y desplazamientos", 0.16, 0)],
    40: [("7.00.22", "Tabla de equivalencias entre UT, UV, distancia y cobertura", 0.16, 0)],
    41: [
        ("7.00.28", "Selector del Modo de mantenimiento", 0.16, 0),
        ("7.00.34", "Control de música y volumen de la cuenta", 0.16, 0),
    ],
}

NOTAS_SIN_CAPTURA = {
    4: "La pantalla de verificación del código no fue capturada; seguí el enlace que llega al correo y sus indicaciones vigentes.",
    5: "La pantalla de nueva contraseña no fue capturada; completá los dos campos que aparezcan después de verificar el código.",
    15: "No hay una captura del panel de capas territoriales abierto. Las imágenes del Modo Libre muestran el acceso al editor y al panel de referencias.",
    18: "No se aportó una captura de Resultados de una ejecución terminada. Consultá allí el detalle de condiciones y puntaje después de simular y finalizar.",
    27: "No se aportó una captura del diálogo Editar usuario; la imagen anterior muestra el botón Editar de cada fila.",
    29: "No se aportó una captura del detalle de un diseño; la imagen anterior muestra cómo abrirlo desde la lista.",
    30: "No se aportó una captura del módulo Actividad. Su lectura se explica a continuación.",
}

def texto_nodo(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))

def encontrar_encabezado(doc, texto):
    return next(p._p for p in doc.paragraphs if p.text == texto and p.style.name == "Heading 1")

def reemplazar_parrafo(parrafo, nuevo):
    if not parrafo:
        return
    for run in parrafo.runs:
        run.text = ""
    parrafo.add_run(nuevo)

def nombre_captura(hora):
    candidatos = list(CAPTURAS.glob(f"Captura de pantalla 2026-10-18 a las {hora}*"))
    assert len(candidatos) == 1, (hora, candidatos)
    return candidatos[0]

def imagen_recortada(doc, ancla, archivo, leyenda, izquierda=0, derecha=0):
    # La franja superior contiene sólo la ventana del navegador (aprox. 13,2 %).
    arriba = 0.132
    ancho_px, alto_px = 2940, 1846
    ancho = 6.02
    alto = ancho * alto_px * (1 - arriba) / (ancho_px * (1 - izquierda - derecha))
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.keep_with_next = True
    parrafo.paragraph_format.space_before = Pt(8)
    parrafo.paragraph_format.space_after = Pt(2)
    forma = parrafo.add_run().add_picture(str(archivo), width=Inches(ancho))
    forma.height = Inches(alto)
    forma._inline.docPr.set("descr", leyenda)
    blip_fill = forma._inline.graphic.graphicData.pic.blipFill
    corte = OxmlElement("a:srcRect")
    corte.set("t", str(round(arriba * 100000)))
    if izquierda:
        corte.set("l", str(round(izquierda * 100000)))
    if derecha:
        corte.set("r", str(round(derecha * 100000)))
    blip_fill.insert(1, corte)
    ancla._p.addprevious(parrafo._p)


d = Document(ORIGEN)
body = d.element.body
# La cubierta usa el logotipo final que aportó la persona usuaria.
logo = next(p for p in d.paragraphs if p._p.xpath(".//a:blip"))
logo.clear()
logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo.add_run().add_picture(str(BASE / "materiales" / "logo METRONET Nuevo.png"), width=Inches(3.4))
body.remove(logo._p)
body.insert(0, logo._p)
for p in list(d.paragraphs):
    if p.text.startswith("Figura 1. Logotipo original"):
        body.remove(p._p)
inicio_usuario = encontrar_encabezado(d, "PARTE A  USUARIO")
inicio_admin = encontrar_encabezado(d, "PARTE B  ADMINISTRADOR")
inicio_problemas = encontrar_encabezado(d, "14 Problemas habituales")

# Reorganizar capítulos: primero Administración y después Usuario.
nodos = list(body)
i_usuario, i_admin, i_problemas = (nodos.index(x) for x in (inicio_usuario, inicio_admin, inicio_problemas))
bloque_admin = nodos[i_admin:i_problemas]
for nodo in bloque_admin:
    body.remove(nodo)
for nodo in bloque_admin:
    inicio_usuario.addprevious(nodo)

# Cubierta y alcance: el manual describe la interfaz vista, con advertencia de rol.
titulo_cubierta = next(p for p in d.paragraphs if p.text == "Manual de usuario METRONET")
subtitulo_cubierta = next(p for p in d.paragraphs if p.text.startswith("Guía práctica para estudiantes"))
reemplazar_parrafo(titulo_cubierta, "Manual de administración y usuario de METRONET")
reemplazar_parrafo(subtitulo_cubierta, "Parte 1 Administración · Parte 2 Usuario\nGuía visual de las pantallas y tareas del juego")
titulo_cubierta.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitulo_cubierta.alignment = WD_ALIGN_PARAGRAPH.CENTER
for par in d.paragraphs:
    if par.text.startswith("Versión local comprobada:"):
        reemplazar_parrafo(par, "Edición del manual: 18 de octubre de 2026.")
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif par.text.startswith("Alcance de la evidencia."):
        reemplazar_parrafo(par, "Cómo usar esta guía. Las capturas muestran la aplicación en una sesión administradora. En la Parte 2, los niveles disponibles y ciertos accesos pueden variar para cada jugador. Las imágenes muestran pantallas observadas; cuando se explica una función sin captura, el texto lo indica. No se modificaron cuentas, partidas ni publicaciones para preparar este manual.")

# Reconstruir el índice en el orden solicitado.
indice = next(p for p in d.paragraphs if p.text == "Índice" and p.style.name == "Heading 1")
primer_parte = next(p for p in d.paragraphs if p.text == "PARTE B  ADMINISTRADOR" and p.style.name == "Heading 1")
cursor = indice._p.getnext()
while cursor is not primer_parte._p:
    siguiente = cursor.getnext()
    body.remove(cursor)
    cursor = siguiente
entradas = [
    "PARTE 1  ADMINISTRADOR",
    "1 Ingreso y navegación de Administración",
    "2 Usuarios",
    "3 Diseños y actividad",
    "4 Experiencia de juego",
    "5 Fuentes, unidades y configuración",
    "PARTE 2  USUARIO",
    "6 Registro, acceso, recuperación y perfil",
    "7 Campaña y niveles",
    "8 Tutoriales y Extras",
    "9 Edición de la red",
    "10 Mapa, capas, objetivos y POI",
    "11 Guardado, simulación y UV y UT",
    "12 Resultados, diseños, aprendizaje y reglas",
    "13 Ranking, trofeos y Modo Libre",
    "14 Problemas habituales",
    "15 Glosario",
]
for linea in entradas:
    nuevo = d.add_paragraph(linea)
    primer_parte._p.addprevious(nuevo._p)

reemplazos = {
    "PARTE B  ADMINISTRADOR": "PARTE 1  ADMINISTRADOR",
    "PARTE A  USUARIO": "PARTE 2  USUARIO",
    "9 Ingreso y navegación de Administración": "1 Ingreso y navegación de Administración",
    "10 Usuarios": "2 Usuarios",
    "11 Diseños y actividad": "3 Diseños y actividad",
    "12 Experiencia de juego": "4 Experiencia de juego",
    "13 Fuentes, unidades y configuración": "5 Fuentes, unidades y configuración",
    "1 Registro, acceso, recuperación y perfil": "6 Registro, acceso, recuperación y perfil",
    "2 Campaña y niveles": "7 Campaña y niveles",
    "3 Tutoriales y Extras": "8 Tutoriales y Extras",
    "4 Edición de la red": "9 Edición de la red",
    "5 Mapa, capas, objetivos y POI": "10 Mapa, capas, objetivos y POI",
    "6 Guardado, simulación y UV y UT": "11 Guardado, simulación y UV y UT",
    "7 Resultados, diseños, aprendizaje y reglas": "12 Resultados, diseños, aprendizaje y reglas",
    "8 Ranking, trofeos y Modo Libre": "13 Ranking, trofeos y Modo Libre",
}
for par in d.paragraphs:
    if par.style.name.startswith("Heading") and par.text in reemplazos:
        reemplazar_parrafo(par, reemplazos[par.text])
    if par.text.startswith("Los capítulos 1 a 8 describen"):
        reemplazar_parrafo(par, "Esta parte describe la experiencia de una cuenta jugadora. Las capturas se tomaron durante una sesión administradora; los niveles disponibles y los accesos del jugador dependen de su progreso y sus permisos.")
    elif par.text.startswith("En la visita al nivel 1"):
        reemplazar_parrafo(par, "En el nivel 1 se observan cinco metas: ubicar dos estaciones, crear una línea, crear una conexión, asignar una unidad y simular la red con sus UV y UT actuales. La vista muestra las herramientas de construcción, edición y acciones. El mapa ofrece zoom y controles de orientación.")
    elif par.text.startswith("Comprobación en vivo: al pulsar Simular diseño"):
        reemplazar_parrafo(par, "La captura muestra el aviso inicial sobre UV y UT al preparar la simulación. No hay una imagen de una ejecución terminada; los pasos siguientes explican los controles visibles.")
    elif par.text.startswith("La observación de esta guía fue de consulta:"):
        reemplazar_parrafo(par, "Las capturas documentan la interfaz administrativa. Los pasos de escritura explican sus controles, pero no implican que se hayan cambiado roles, borradores, publicaciones, diseños o ajustes en esta sesión.")
    elif par.text.startswith("Hay una contradicción visible:"):
        reemplazar_parrafo(par, "Nota sobre las reglas: esta pantalla indica que UV y UT son independientes en Niveles, mientras que el recorrido actual las relaciona por la distancia de cada línea. Al ajustar una línea, observá el valor recalculado antes de simular. Esta diferencia de información queda señalada para revisión del proyecto.")
    elif "con datos personales ocultados" in par.text or "con nombres y correos ocultados" in par.text:
        pass

# Reemplazar cada marcador de captura por figuras actuales o por una nota honesta.
usadas = set()
contador = 0
for par in list(d.paragraphs):
    match = re.match(r"Captura pendiente (\d+)\. ", par.text)
    if not match:
        continue
    numero = int(match.group(1))
    if numero in FIGURAS:
        for hora, leyenda, izquierda, derecha in FIGURAS[numero]:
            archivo = nombre_captura(hora)
            usadas.add(archivo.name)
            contador += 1
            imagen_recortada(d, par, archivo, leyenda, izquierda, derecha)
            etiqueta = d.add_paragraph(f"Figura {contador}. {leyenda}.", style="Caption")
            etiqueta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            etiqueta.paragraph_format.keep_with_next = False
            par._p.addprevious(etiqueta._p)
        par._p.getparent().remove(par._p)
    else:
        reemplazar_parrafo(par, NOTAS_SIN_CAPTURA[numero])
        par.style = d.styles["Normal"]

# Ajustes finales de texto y paginación.
for par in d.paragraphs:
    if par.text.startswith("Esta parte corresponde exclusivamente al rol administrador"):
        par.text = "Esta parte corresponde al rol administrador. Una cuenta jugadora no ve el enlace Administración ni sus módulos. Las pantallas generales pueden consultarse desde ambos roles, pero el administrador no recibe marcas ni trofeos de jugador."
    if "Los lugares de las capturas" in par.text:
        raise AssertionError(par.text)
    if par.text.startswith("En Mis diseños, cada tarjeta"):
        par.text = par.text.replace("Los contadores de las pestañas indican cuántos diseños hay en cada grupo.", "Las pestañas separan los diseños de niveles y los de Modo Libre.")
    if par.text.startswith("Para el estudiante, la navegación ofrece"):
        par.text = par.text.replace("Durante un nivel se indica la secuencia Nivel → Edición → Simulación → Resultados.", "Durante un nivel seguí la secuencia Nivel → Edición → Simulación → Resultados.")
    if "Abrí la dirección local que usa el proyecto" in par.text:
        par.text = par.text.replace("Abrí la dirección local que usa el proyecto, normalmente http://127.0.0.1:5173/login.html, con frontend y backend ya iniciados.", "Abrí la dirección de METRONET que te haya facilitado el equipo responsable.")

# Apariencia de manual visual basada en el ejemplo aportado: blanco, azul y tablas claras.
for estilo in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet"):
    st = d.styles[estilo]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
d.styles["Normal"].font.size = Pt(9)
d.styles["Normal"].paragraph_format.space_after = Pt(4)
d.styles["Normal"].paragraph_format.line_spacing = 1.1
d.styles["Title"].font.size = Pt(20)
d.styles["Title"].font.color.rgb = RGBColor(16, 89, 171)
d.styles["Title"].paragraph_format.space_after = Pt(12)
d.styles["Heading 1"].font.size = Pt(14)
d.styles["Heading 1"].font.color.rgb = RGBColor(16, 89, 171)
d.styles["Heading 2"].font.size = Pt(10.5)
d.styles["Heading 2"].font.color.rgb = RGBColor(33, 43, 56)
d.styles["Caption"].font.size = Pt(7.5)
d.styles["Caption"].font.bold = False
d.styles["Caption"].font.italic = True
d.styles["Caption"].font.color.rgb = RGBColor(91, 104, 119)
d.styles["Caption"].paragraph_format.space_after = Pt(5)
for seccion in d.sections:
    seccion.page_width = Inches(8.27)
    seccion.page_height = Inches(11.69)
for tabla in d.tables:
    for fila in tabla.rows:
        if fila.cells[0].text.startswith("No abre la aplicación local"):
            fila.cells[0].text = "No abre METRONET"
            fila.cells[1].text = "Comprobá la dirección facilitada y consultá al equipo responsable del acceso."
    for ri, fila in enumerate(tabla.rows):
        for celda in fila.cells:
            tcPr = celda._tc.get_or_add_tcPr()
            for sh in tcPr.findall(qn("w:shd")):
                tcPr.remove(sh)
            sh = OxmlElement("w:shd")
            sh.set(qn("w:fill"), "EAF2FA" if ri == 0 else ("F7FAFD" if ri % 2 == 0 else "FFFFFF"))
            tcPr.append(sh)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(31, 42, 55)
                    if ri == 0:
                        r.bold = True

pie = d.sections[0].footer.paragraphs[0]
pie.text = "METRONET · Manual integral · 18.10.2026 · "
campo = OxmlElement("w:fldSimple")
campo.set(qn("w:instr"), "PAGE")
pie._p.append(campo)
pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in pie.runs:
    r.font.size = Pt(8)

d.core_properties.title = "Manual integral METRONET"
d.core_properties.subject = "Parte 1 Administración y Parte 2 Usuario"
d.core_properties.author = "METRONET"

originales = {p.name for p in CAPTURAS.glob("Captura de pantalla 2026-10-18*.png")}
assert len(originales) == 46, len(originales)
assert usadas == originales, f"Sin usar: {sorted(originales-usadas)}; desconocidas: {sorted(usadas-originales)}"
assert not any(p.text.startswith("Captura pendiente") for p in d.paragraphs)
d.save(SALIDA)
print(f"Documento: {SALIDA}\nCapturas originales incluidas: {len(usadas)}\nFiguras: {contador}")
